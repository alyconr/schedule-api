import io
import json
import re
import hashlib
import unicodedata
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any, Literal, Optional
import openpyxl
from sqlmodel import Session, select

from app.models.academic import (
    ContractType,
    Instructor,
    Environment,
    TrainingProgram,
    Group,
    Competency,
    LearningResult,
    Topic,
    LearningResultTopic,
    AcademicPeriod,
)
from app.models.auth import User
from app.models.coordination import Coordination, InstructorCoordination, EnvironmentCoordination
from app.models.imports import ImportBatch, ImportBatchRecord
from app.services.imports import (
    WorkbookProfileDetector,
    parse_instructors_sheet,
    parse_environments_sheet,
    parse_groups_sheet,
    parse_academic_periods_sheet,
    parse_semaforo_sheet as parse_canonical_semaforo_sheet,
    canonical_program_key,
    build_program_code as canonical_build_program_code,
)
from app.services.imports.safe_merge import (
    compute_file_sha256,
    compute_record_fingerprint,
    merge_field,
    deduplicate_records,
    ClassifiedRecord,
    classify_group,
    classify_instructor,
    classify_environment,
    classify_training_program,
    classify_academic_period,
    classify_contract_type,
    classify_generic_global_entity,
)
from app.schemas.imports import (
    ImportIssue,
    ImportEntitySummary,
    ImportPreviewResponse,
    ImportCommitResponse,
    ImportConflictItem,
    ImportChangeItem,
    ImportFieldChange,
)



def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    s = str(value).strip().lower()
    s = "".join(
        c for c in unicodedata.normalize("NFD", s)
        if unicodedata.category(c) != "Mn"
    )
    s = re.sub(r"[\s\-]+", "_", s)
    s = re.sub(r"[^\w_]", "", s)
    return s


def normalize_program_name(value: Any) -> str:
    return " ".join(str(value or "").strip().upper().strip(" .;:,").split())


def program_code_from_name(value: Any) -> str:
    return canonical_build_program_code(value)


def build_program_code(program_name: str) -> str:
    return canonical_build_program_code(program_name)



def parse_excel_date(value: Any) -> Optional[date]:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        if isinstance(value, datetime):
            return value.date()
        return value
    if isinstance(value, (int, float)):
        try:
            return date(1899, 12, 30) + timedelta(days=float(value))
        except Exception:
            pass
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def get_stable_hash(text: str) -> str:
    return hashlib.md5(text.encode("utf-8")).hexdigest()[:8].upper()


def find_col_idx(headers: list[str], *keys: str) -> Optional[int]:
    for key in keys:
        for i, h in enumerate(headers):
            if h == key:
                return i
    for key in keys:
        for i, h in enumerate(headers):
            if h.startswith(key):
                return i
    for key in keys:
        for i, h in enumerate(headers):
            if key in h:
                return i
    return None


ROMAN_TRIMESTERS = {"i": 1, "ii": 2, "iii": 3, "iv": 4, "v": 5, "vi": 6}
NORMALIZED_SCHEDULE_SHEETS = {
    "LISTA INSTRUCTORES": "lista_instructores",
    "AMBIENTES": "ambientes",
    "FICHAS": "fichas",
    "Semaforo con RA cadena": "semaforo_con_ra_cadena",
    "Semaforo con RA Oferta Abierta": "semaforo_con_ra_oferta_abierta",
}


def normalize_contract_type(value: Any) -> dict[str, Any]:
    raw = normalize_header(value)
    if raw in (
        "carrera_administrativa",
        "nombramiento_ordinario",
        "nombramiento_provisional",
        "planta",
        "nombramiento",
        "empleado_publico",
    ):
        return {"name": "planta", "base_hours": Decimal("40"), "max_hours": Decimal("40"), "known": True}
    if raw in ("contratista_sena", "contrato", "contratista", "prestacion_de_servicios", "prestacion_servicios", "contrato_sena"):
        return {"name": "contratista", "base_hours": Decimal("40"), "max_hours": Decimal("48"), "known": True}
    if raw == "otro":
        return {"name": "otro", "base_hours": Decimal("0"), "max_hours": Decimal("0"), "known": True}
    return {"name": "otro", "base_hours": Decimal("0"), "max_hours": Decimal("0"), "known": False}


def normalize_vinculation_label(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().upper().split())


def vinculation_category(value: Any) -> str:
    raw = normalize_header(value)
    if raw in (
        "carrera_administrativa",
        "nombramiento_ordinario",
        "nombramiento_provisional",
        "planta",
    ):
        return "planta"
    if raw in (
        "contratista_sena",
        "contratista",
        "contrato",
        "prestacion_de_servicios",
        "prestacion_servicios",
    ):
        return "contratista"
    return "otro"


def weekly_rule_hours_for_vinculation(value: Any) -> Decimal:
    category = vinculation_category(value)
    if category == "planta":
        return Decimal("32")
    if category == "contratista":
        return Decimal("40")
    return Decimal("0")


def parse_decimal(value: Any) -> Optional[Decimal]:
    if value is None or value == "":
        return None
    try:
        return Decimal(str(value).replace(",", "."))
    except Exception:
        return None


def parse_trimester_label(value: Any) -> tuple[str | None, int | None]:
    if value is None:
        return None, None
    raw = str(value).strip()
    if not raw:
        return None, None
    text = normalize_header(raw).replace("_", " ")
    match = re.search(r"(?:trimestre|trim|t)?\s*([0-9]+|[ivx]+)(?:\s*(?:a|-|al)\s*([0-9]+|[ivx]+))?", text)
    if not match:
        return raw, None
    token = match.group(1)
    number = int(token) if token.isdigit() else ROMAN_TRIMESTERS.get(token)
    return raw, number


def normalized_text(value: Any) -> str:
    return normalize_header(value).replace("_", ".")


def sheet_by_normalized_name(wb: Any, expected: str) -> Any | None:
    expected_key = normalize_header(expected)
    for name in wb.sheetnames:
        if normalize_header(name) == expected_key:
            return wb[name]
    return None


def row_cell(row: tuple[Any, ...], idx: int | None) -> Any:
    return row[idx] if idx is not None and idx < len(row) else None


def add_missing_columns_errors(sheet_name: str, headers: list[str], columns: dict[str, tuple[str, ...]], errors: list[ImportIssue]) -> bool:
    missing = False
    for label, keys in columns.items():
        if find_col_idx(headers, *keys) is None:
            errors.append(ImportIssue(
                sheet=sheet_name,
                severity="error",
                message=f"Falta la columna obligatoria '{label}'.",
            ))
            missing = True
    return missing


def get_cell_fill_key(cell: Any) -> Optional[str]:
    fill = getattr(cell, "fill", None)
    color = getattr(fill, "fgColor", None) or getattr(fill, "start_color", None)
    if not color or getattr(color, "type", None) is None:
        return None
    color_type = color.type
    if color_type == "rgb" and color.rgb and color.rgb not in ("00000000", "FFFFFFFF"):
        return f"rgb:{color.rgb}"
    if color_type == "theme" and color.theme is not None:
        return f"theme:{color.theme}:tint:{round(float(color.tint or 0), 10)}"
    if color_type == "indexed" and color.indexed not in (None, 64):
        return f"indexed:{color.indexed}:tint:{round(float(color.tint or 0), 10)}"
    return None


def get_cell_fill_hex(cell: Any) -> Optional[str]:
    color = getattr(getattr(cell, "fill", None), "fgColor", None)
    if color and color.type == "rgb" and color.rgb and len(color.rgb) >= 6:
        return color.rgb[-6:]
    return None


def detect_trimester(value: Any) -> tuple[Optional[str], Optional[int]]:
    if value is None:
        return None, None
    raw = str(value).strip()
    text = normalize_header(raw).replace("_", " ")
    match = re.search(r"(?:trimestre|trim|t)\s*([0-9]+|[ivx]+)", text)
    if not match:
        match = re.search(r"\b([ivx]+)\s*trimestre\b", text)
    if not match:
        return None, None
    token = match.group(1)
    number = int(token) if token.isdigit() else ROMAN_TRIMESTERS.get(token)
    if number is None:
        return None, None
    return f"TRIMESTRE {number}", number


def clean_learning_text(value: str) -> tuple[str, Optional[int]]:
    clean = " ".join(value.strip().split())
    match = re.match(r"^(\d+)[\.\-\)]?\s*(.*)$", clean)
    if match and match.group(2).strip():
        return match.group(2).strip(), int(match.group(1))
    return clean, None


def semaforo_context_sheets(wb: Any, context: str) -> tuple[Optional[str], Optional[str]]:
    context_key = "oferta_abierta" if context == "oferta_abierta" else "cadena"
    ra_sheet = None
    topic_sheet = None
    for name in wb.sheetnames:
        n = normalize_header(name)
        if context_key not in n:
            continue
        is_topic = "tematic" in n or "tema" in n
        if is_topic and topic_sheet is None:
            topic_sheet = name
        elif ("ra" in n or "semaforo" in n) and ra_sheet is None:
            ra_sheet = name
    return ra_sheet, topic_sheet


def parse_relational_semaforo_sheet(
    sheet: Any,
    context: str,
    kind: Literal["ra", "topic"],
    warnings: list[ImportIssue],
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    current_trimester = None
    current_trimester_number = None
    order_by_trimester: dict[int, int] = {}

    for row in sheet.iter_rows():
        detected = next((detect_trimester(cell.value) for cell in row if detect_trimester(cell.value)[1]), (None, None))
        if detected[1]:
            current_trimester, current_trimester_number = detected
            continue

        for idx, cell in enumerate(row[:-1]):
            value = cell.value
            hours_cell = row[idx + 1]
            hours = parse_decimal(hours_cell.value)
            if not isinstance(value, str) or not value.strip() or hours is None:
                continue
            raw_text = value.strip()
            header = normalize_header(raw_text)
            if header.startswith(("trimestre", "actividad", "horas", "hora", "2025", "2026")):
                continue

            color_key = get_cell_fill_key(cell)
            color_hex = get_cell_fill_hex(cell)
            if current_trimester_number is None:
                warnings.append(ImportIssue(
                    sheet=sheet.title,
                    row=cell.row,
                    entity=kind,
                    severity="warning",
                    message="Fila sin trimestre detectable; no se usara para relacionar por color.",
                    raw_value=raw_text,
                ))
            if not color_key:
                warnings.append(ImportIssue(
                    sheet=sheet.title,
                    row=cell.row,
                    entity=kind,
                    severity="warning",
                    message="Fila sin color de fondo util; no se usara para relacionar por color.",
                    raw_value=raw_text,
                ))

            tri_num = current_trimester_number or 0
            order_by_trimester[tri_num] = order_by_trimester.get(tri_num, 0) + 1
            prefix = "OA" if context == "oferta_abierta" else "CAD"
            code_kind = "RA" if kind == "ra" else "TEM"
            code = f"{prefix}_T{tri_num:02d}_{code_kind}_{order_by_trimester[tri_num]:03d}"
            description, number_prefix = clean_learning_text(raw_text)
            items.append({
                "code": code,
                "program_scope": context,
                "trimester": current_trimester,
                "trimester_number": current_trimester_number,
                "source_sheet": sheet.title,
                "source_address": cell.coordinate,
                "source_row": cell.row,
                "source_col": cell.column,
                "order_in_trimester": order_by_trimester[tri_num],
                "color_key": color_key,
                "color_hex": color_hex,
                "number_prefix": number_prefix,
                "text_original": raw_text,
                "description": description[:500],
                "estimated_hours": hours,
                "result_type": "semaforo_relacional" if kind == "ra" else "tematica",
            })
    return items


def process_relational_semaforos(wb: Any, warnings: list[ImportIssue]) -> dict[str, list[dict[str, Any]]]:
    learning_results: list[dict[str, Any]] = []
    topics: list[dict[str, Any]] = []
    color_groups: list[dict[str, Any]] = []
    relations: list[dict[str, Any]] = []

    for context in ("oferta_abierta", "cadena"):
        ra_sheet_name, topic_sheet_name = semaforo_context_sheets(wb, context)
        if not ra_sheet_name or not topic_sheet_name:
            continue
        context_ras = parse_relational_semaforo_sheet(wb[ra_sheet_name], context, "ra", warnings)
        context_topics = parse_relational_semaforo_sheet(wb[topic_sheet_name], context, "topic", warnings)
        learning_results.extend(context_ras)
        topics.extend(context_topics)

        group_keys = {
            (item["program_scope"], item["trimester_number"], item["color_key"])
            for item in context_ras + context_topics
            if item.get("trimester_number") and item.get("color_key")
        }
        for group_index, key in enumerate(sorted(group_keys, key=lambda k: (k[0], k[1], k[2])), start=1):
            scope, trimester_number, color_key = key
            group_ras = [item for item in context_ras if (item["program_scope"], item["trimester_number"], item["color_key"]) == key]
            group_topics = [item for item in context_topics if (item["program_scope"], item["trimester_number"], item["color_key"]) == key]
            prefix = "OA" if scope == "oferta_abierta" else "CAD"
            group_id = f"{prefix}_T{trimester_number:02d}_COLOR_{group_index:03d}"
            status = "BLOQUE_MUCHOS_A_MUCHOS" if len(group_ras) > 1 and len(group_topics) > 1 else "OK"
            needs_review = status != "OK"
            if not group_ras or not group_topics:
                warnings.append(ImportIssue(
                    sheet="semaforos_relacional",
                    entity="color_group",
                    severity="warning",
                    message="Grupo de color sin pareja RA/tematica.",
                    raw_value=f"{scope} T{trimester_number} {color_key}",
                ))
                continue
            color_groups.append({
                "group_id": group_id,
                "program_scope": scope,
                "trimester_number": trimester_number,
                "color_key": color_key,
                "color_hex": next((item.get("color_hex") for item in group_ras + group_topics if item.get("color_hex")), None),
                "ra_count": len(group_ras),
                "topic_count": len(group_topics),
                "relation_status": status,
                "confidence": "media" if needs_review else "alta",
                "needs_manual_review": needs_review,
            })
            for ra in group_ras:
                for topic in group_topics:
                    relations.append({
                        "relation_id": f"{group_id}_{ra['code']}_{topic['code']}"[:120],
                        "group_id": group_id,
                        "learning_result_code": ra["code"],
                        "topic_code": topic["code"],
                        "program_scope": scope,
                        "trimester_number": trimester_number,
                        "color_key": color_key,
                        "color_hex": next((item.get("color_hex") for item in (ra, topic) if item.get("color_hex")), None),
                        "relation_method": "same_trimester_and_same_fill_color",
                        "relation_status": status,
                        "confidence": "media" if needs_review else "alta",
                        "needs_manual_review": needs_review,
                    })

    return {
        "learning_results": learning_results,
        "topics": topics,
        "color_groups": color_groups,
        "ra_topic_relations": relations,
    }


def sheet_records(sheet: Any) -> list[dict[str, Any]]:
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [normalize_header(cell) for cell in rows[0]]
    records = []
    for row in rows[1:]:
        if not any(cell is not None for cell in row):
            continue
        records.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers))})
    return records


def bool_value(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in ("true", "1", "si", "sí", "yes")


def process_normalized_relational_workbook(wb: Any) -> dict[str, list[dict[str, Any]]]:
    learning_results = []
    topics = []
    color_groups = []
    relations = []

    for record in sheet_records(wb["learning_results"]):
        learning_results.append({
            "code": str(record["learning_result_id"])[:50],
            "program_scope": normalize_header(record.get("program_scope")),
            "trimester": record.get("trimester"),
            "trimester_number": int(record["trimester_number"]) if record.get("trimester_number") is not None else None,
            "source_sheet": record.get("source_sheet"),
            "source_address": record.get("source_address"),
            "source_row": record.get("source_row"),
            "source_col": record.get("source_col"),
            "order_in_trimester": record.get("order_in_trimester"),
            "color_key": record.get("color_key"),
            "color_hex": record.get("color_hex"),
            "number_prefix": record.get("number_prefix"),
            "text_original": record.get("text_original"),
            "description": str(record.get("description_clean") or record.get("text_original") or "")[:500],
            "estimated_hours": parse_decimal(record.get("hours")),
            "result_type": "semaforo_relacional",
        })

    for record in sheet_records(wb["topics"]):
        topics.append({
            "code": str(record["topic_id"])[:50],
            "program_scope": normalize_header(record.get("program_scope")),
            "trimester": record.get("trimester"),
            "trimester_number": int(record["trimester_number"]) if record.get("trimester_number") is not None else None,
            "source_sheet": record.get("source_sheet"),
            "source_address": record.get("source_address"),
            "source_row": record.get("source_row"),
            "source_col": record.get("source_col"),
            "order_in_trimester": record.get("order_in_trimester"),
            "color_key": record.get("color_key"),
            "color_hex": record.get("color_hex"),
            "text_original": record.get("text_original"),
            "description": str(record.get("description_clean") or record.get("text_original") or "")[:500],
            "estimated_hours": parse_decimal(record.get("hours")),
            "result_type": "tematica",
        })

    for record in sheet_records(wb["color_groups"]):
        color_groups.append({
            "group_id": record.get("group_id"),
            "program_scope": normalize_header(record.get("program_scope")),
            "trimester_number": int(record["trimester_number"]) if record.get("trimester_number") is not None else None,
            "color_key": record.get("color_key"),
            "color_hex": record.get("color_hex"),
            "ra_count": record.get("ra_count"),
            "topic_count": record.get("topic_count"),
            "relation_status": record.get("relation_status"),
            "confidence": record.get("confidence"),
            "needs_manual_review": record.get("relation_status") == "BLOQUE_MUCHOS_A_MUCHOS",
        })

    for record in sheet_records(wb["ra_topic_relations"]):
        relations.append({
            "relation_id": str(record["relation_id"])[:120],
            "group_id": record.get("group_id"),
            "learning_result_code": str(record["learning_result_id"])[:50],
            "topic_code": str(record["topic_id"])[:50],
            "program_scope": normalize_header(record.get("program_scope")),
            "trimester_number": int(record["trimester_number"]) if record.get("trimester_number") is not None else None,
            "color_key": record.get("color_key"),
            "color_hex": record.get("color_hex"),
            "relation_method": record.get("relation_method") or "same_trimester_and_same_fill_color",
            "relation_status": record.get("relation_status") or "OK",
            "confidence": record.get("confidence") or "alta",
            "needs_manual_review": bool_value(record.get("needs_manual_review")),
        })

    return {
        "learning_results": learning_results,
        "topics": topics,
        "color_groups": color_groups,
        "ra_topic_relations": relations,
    }


def process_schedule_normalized_workbook(wb: Any, warnings: list[ImportIssue], errors: list[ImportIssue]) -> dict[str, list[dict[str, Any]]]:
    items: dict[str, list[dict[str, Any]]] = {
        "academic_periods": [],
        "contract_types": [],
        "instructors": [],
        "environments": [],
        "groups": [],
        "programs": [],
        "learning_results": [],
        "topics": [],
        "color_groups": [],
        "ra_topic_relations": [],
    }

    resolved = WorkbookProfileDetector.resolve(wb, errors)
    if errors:
        return items

    # 1. Parse Academic Periods (if sheet exists)
    periods_sheet = resolved.sheets.get("academic_periods")
    if periods_sheet is not None:
        items["academic_periods"] = parse_academic_periods_sheet(
            periods_sheet,
            parse_excel_date,
            warnings,
            errors,
        )

    # 2. Parse Instructors & Contract Types
    inst_sheet = resolved.sheets.get("instructors")
    if inst_sheet is not None:
        insts, c_types = parse_instructors_sheet(
            inst_sheet,
            normalize_contract_type,
            vinculation_category,
            weekly_rule_hours_for_vinculation,
            warnings,
            errors,
        )
        items["instructors"] = insts
        items["contract_types"] = c_types

    # 3. Parse Environments
    env_sheet = resolved.sheets.get("environments")
    if env_sheet is not None:
        items["environments"] = parse_environments_sheet(env_sheet, warnings, errors)

    # 4. Parse Groups (FICHAS) & extract program scope map
    groups_sheet = resolved.sheets.get("groups")
    program_scope_map: dict[str, str] = {}
    if groups_sheet is not None:
        grps, progs, scope_map = parse_groups_sheet(
            groups_sheet,
            parse_excel_date,
            warnings,
            errors,
        )
        items["groups"] = grps
        items["programs"] = progs
        program_scope_map = scope_map

    seen_programs = {p["code"] for p in items["programs"]}
    seen_ras: set[str] = set()
    seen_topics: set[str] = set()
    seen_relations: set[str] = set()

    # 5. Parse Semaforo sheets with the generic parser
    for sem_sheet, sem_config in resolved.semaforo_sheets:
        parse_canonical_semaforo_sheet(
            sheet=sem_sheet,
            config=sem_config,
            program_scope_map=program_scope_map,
            programs_list=items["programs"],
            learning_results_list=items["learning_results"],
            topics_list=items["topics"],
            relations_list=items["ra_topic_relations"],
            seen_programs=seen_programs,
            seen_ras=seen_ras,
            seen_topics=seen_topics,
            seen_relations=seen_relations,
            warnings=warnings,
            errors=errors,
        )

    return items



def parse_semaforo_sheet(sheet: Any) -> list[dict[str, Any]]:
    items = []
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < 2:
        return items
    
    for r_idx in range(1, len(rows)):
        row = rows[r_idx]
        for c_idx in range(len(row) - 1):
            val = row[c_idx]
            hours_val = row[c_idx + 1]
            if isinstance(val, str) and val.strip():
                clean_val = val.strip()
                if clean_val.lower().startswith(('trimestre', 'actividad', 'horas', 'hora', '2025', '2026')):
                    continue
                if isinstance(hours_val, (int, float)):
                    items.append({
                        "description": clean_val,
                        "hours": float(hours_val)
                    })
    return items


def upsert_entity(session: Session, model_cls: Any, lookup_field: str, lookup_value: Any, payload: dict) -> tuple[Any, bool]:
    stmt = select(model_cls).where(getattr(model_cls, lookup_field) == lookup_value)
    existing = session.exec(stmt).first()
    if existing:
        for k, v in payload.items():
            merged = merge_field(getattr(existing, k, None), v)
            setattr(existing, k, merged)
        session.add(existing)
        session.flush()
        return existing, False
    else:
        new_obj = model_cls(**payload)
        session.add(new_obj)
        session.flush()
        return new_obj, True



def process_workbook(file_bytes: bytes, filename: str, import_type: str) -> dict[str, Any]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheets_detected = wb.sheetnames
    
    warnings: list[ImportIssue] = []
    errors: list[ImportIssue] = []
    
    instructors_list = []
    environments_list = []
    groups_list = []
    programs_list = []
    learning_results_list = []
    topics_list = []
    color_groups_list = []
    ra_topic_relations_list = []
    
    counts = {
        "academic_periods": {"valid": 0, "warnings": 0, "rejected": 0},
        "contract_types": {"valid": 0, "warnings": 0, "rejected": 0},
        "instructors": {"valid": 0, "warnings": 0, "rejected": 0},
        "environments": {"valid": 0, "warnings": 0, "rejected": 0},
        "groups": {"valid": 0, "warnings": 0, "rejected": 0},
        "programs": {"valid": 0, "warnings": 0, "rejected": 0},
        "learning_results": {"valid": 0, "warnings": 0, "rejected": 0},
        "topics": {"valid": 0, "warnings": 0, "rejected": 0},
        "color_groups": {"valid": 0, "warnings": 0, "rejected": 0},
        "ra_topic_relations": {"valid": 0, "warnings": 0, "rejected": 0}
    }

    if import_type == "schedule_normalized":
        normalized = process_schedule_normalized_workbook(wb, warnings, errors)
        if not errors:
            normalized["instructors"] = deduplicate_records(
                normalized.get("instructors", []),
                lambda x: x.get("document_number", ""),
                "instructors",
                "LISTA INSTRUCTORES",
                warnings,
                errors,
            )
            normalized["environments"] = deduplicate_records(
                normalized.get("environments", []),
                lambda x: x.get("code", ""),
                "environments",
                "AMBIENTES",
                warnings,
                errors,
            )
            normalized["groups"] = deduplicate_records(
                normalized.get("groups", []),
                lambda x: x.get("code", ""),
                "groups",
                "FICHAS",
                warnings,
                errors,
            )
            normalized["programs"] = deduplicate_records(
                normalized.get("programs", []),
                lambda x: x.get("code", ""),
                "programs",
                "FICHAS",
                warnings,
                errors,
            )
        for key, values in normalized.items():
            if key in counts:
                counts[key]["valid"] = len(values)

        for issue in warnings:
            entity_key = {
                "academic_period": "academic_periods",
                "instructor": "instructors",
                "contract_type": "contract_types",
                "environment": "environments",
                "group": "groups",
                "program": "programs",
                "learning_result": "learning_results",
                "topic": "topics",
                "ra_topic_relation": "ra_topic_relations",
            }.get(issue.entity or "")
            if entity_key in counts:
                counts[entity_key]["warnings"] += 1
        for issue in errors:
            entity_key = {
                "academic_period": "academic_periods",
                "instructor": "instructors",
                "contract_type": "contract_types",
                "environment": "environments",
                "group": "groups",
                "program": "programs",
                "learning_result": "learning_results",
                "topic": "topics",
                "ra_topic_relation": "ra_topic_relations",
            }.get(issue.entity or "")
            if entity_key in counts:
                counts[entity_key]["rejected"] += 1
        return {
            "import_type": import_type,
            "filename": filename,
            "sheets_detected": sheets_detected,
            "summary": {k: ImportEntitySummary(**v) for k, v in counts.items()},
            "items": normalized,
            "warnings": warnings,
            "errors": errors,
        }
    
    # 1. Parse LISTA_INSTRUCTORES_AMBIENTES
    sheet_ia = "LISTA_INSTRUCTORES_AMBIENTES"
    if import_type in ("semaforos_sena", "instructors_environments"):
        if sheet_ia not in sheets_detected:
            errors.append(ImportIssue(
                sheet=sheet_ia,
                severity="error",
                message=f"La hoja obligatoria '{sheet_ia}' no fue encontrada en el archivo."
            ))
        else:
            sheet = wb[sheet_ia]
            rows = list(sheet.iter_rows(values_only=True))
            if len(rows) < 2:
                errors.append(ImportIssue(
                    sheet=sheet_ia,
                    severity="error",
                    message=f"La hoja '{sheet_ia}' está vacía o no contiene filas de datos."
                ))
            else:
                headers = [normalize_header(c) for c in rows[0]]
                
                # find indices
                nombre_idx = find_col_idx(headers, 'nombre_completo', 'nombre')
                iniciales_idx = find_col_idx(headers, 'iniciales', 'inicial')
                tipo_contrato_idx = find_col_idx(headers, 'tipo_contrato', 'contrato')
                horas_formacion_idx = find_col_idx(headers, 'horasfromacion', 'horas_formacion', 'horas_base')
                horas_adicionales_idx = find_col_idx(headers, 'horas_adicionales', 'adicionales')
                coordinacion_idx = find_col_idx(headers, 'coordinacion')
                numero_idx = find_col_idx(headers, 'numero', 'num')
                ambiente_idx = find_col_idx(headers, 'ambiente', 'nombre_ambiente')
                actividad_idx = find_col_idx(headers, 'actividad_de_formacion', 'actividad')
                horas_actividad_idx = find_col_idx(headers, 'horas')
                
                # Walk rows
                for r_idx in range(1, len(rows)):
                    row = rows[r_idx]
                    row_num = r_idx + 1
                    
                    # check if row is empty
                    if not any(c is not None for c in row):
                        continue
                        
                    has_instructor_err = False
                    has_instructor_warn = False
                    
                    # Parse Instructor
                    nombre = row[nombre_idx] if nombre_idx is not None else None
                    iniciales = str(row[iniciales_idx]).strip() if (iniciales_idx is not None and row[iniciales_idx] is not None) else ""
                    tipo_contrato = str(row[tipo_contrato_idx]).strip() if (tipo_contrato_idx is not None and row[tipo_contrato_idx] is not None) else ""
                    horas_formacion = row[horas_formacion_idx] if horas_formacion_idx is not None else None
                    horas_adicionales = row[horas_adicionales_idx] if horas_adicionales_idx is not None else None
                    coordinacion = str(row[coordinacion_idx]).strip() if (coordinacion_idx is not None and row[coordinacion_idx] is not None) else ""
                    
                    if nombre and str(nombre).strip():
                        nombre_str = str(nombre).strip()
                        parts = nombre_str.split()
                        first_name = parts[0] if len(parts) == 1 else " ".join(parts[:-1])
                        last_name = parts[-1] if len(parts) > 1 else "PENDIENTE"
                        
                        # Document number
                        if iniciales:
                            doc_num = f"TEMP-{iniciales.upper()}"
                        else:
                            doc_num = f"TEMP-{get_stable_hash(nombre_str)}"
                            warnings.append(ImportIssue(
                                sheet=sheet_ia,
                                row=row_num,
                                entity="instructor",
                                severity="warning",
                                message=f"Instructor sin documento. Se generó documento temporal: {doc_num}",
                                raw_value=nombre_str
                            ))
                            has_instructor_warn = True
                            
                        # Email
                        if iniciales:
                            email = f"{iniciales.lower()}@pendiente.sena.local"
                        else:
                            email = f"{first_name.lower().replace(' ', '')}.{last_name.lower().replace(' ', '')}@pendiente.sena.local"
                            warnings.append(ImportIssue(
                                sheet=sheet_ia,
                                row=row_num,
                                entity="instructor",
                                severity="warning",
                                message=f"Instructor sin correo. Se generó correo temporal: {email}",
                                raw_value=nombre_str
                            ))
                            has_instructor_warn = True
                            
                        contract_info = normalize_contract_type(tipo_contrato)
                        if tipo_contrato and not contract_info["known"]:
                            warnings.append(ImportIssue(
                                sheet=sheet_ia,
                                row=row_num,
                                entity="instructor",
                                severity="warning",
                                message="Tipo de contrato no reconocido. Se normalizo como 'otro'.",
                                raw_value=tipo_contrato
                            ))
                            has_instructor_warn = True

                        base_h = parse_decimal(horas_formacion)
                        if base_h is None:
                            base_h = contract_info["base_hours"]
                            if horas_formacion is not None:
                                warnings.append(ImportIssue(
                                    sheet=sheet_ia,
                                    row=row_num,
                                    entity="instructor",
                                    severity="warning",
                                    message=f"Horas de formacion invalidas. Se uso valor por defecto: {base_h}",
                                    raw_value=str(horas_formacion)
                                ))
                                has_instructor_warn = True

                        add_h = parse_decimal(horas_adicionales)
                        if add_h is None:
                            max_h = contract_info["max_hours"] if horas_adicionales is None else base_h
                            if horas_adicionales is not None:
                                warnings.append(ImportIssue(
                                    sheet=sheet_ia,
                                    row=row_num,
                                    entity="instructor",
                                    severity="warning",
                                    message="Horas adicionales invalidas. Se uso la carga base como maximo.",
                                    raw_value=str(horas_adicionales)
                                ))
                                has_instructor_warn = True
                        else:
                            max_h = base_h + add_h
                            
                        instructor_dict = {
                            "document_type": "CC",
                            "document_number": doc_num,
                            "first_name": first_name[:100],
                            "last_name": last_name[:100],
                            "email": email[:200],
                            "weekly_base_hours": base_h,
                            "weekly_max_hours": max_h,
                            "area": coordinacion[:100] if coordinacion else None,
                            "contract_type_name": contract_info["name"],
                            "contract_type_base_hours": contract_info["base_hours"],
                            "contract_type_max_hours": contract_info["max_hours"],
                        }
                        instructors_list.append(instructor_dict)
                        
                        if has_instructor_warn:
                            counts["instructors"]["warnings"] += 1
                        else:
                            counts["instructors"]["valid"] += 1
                        
                    # Parse Environment
                    numero = row[numero_idx] if numero_idx is not None else None
                    ambiente = row[ambiente_idx] if ambiente_idx is not None else None
                    
                    has_env_err = False
                    has_env_warn = False
                    
                    if numero is not None or ambiente is not None:
                        code = str(numero).strip() if numero is not None else ""
                        name = str(ambiente).strip() if ambiente is not None else ""
                        
                        if not code and name:
                            code = f"ENV-{get_stable_hash(name)}"
                            warnings.append(ImportIssue(
                                sheet=sheet_ia,
                                row=row_num,
                                entity="environment",
                                severity="warning",
                                message=f"Ambiente sin número. Se generó código temporal: {code}",
                                raw_value=name
                            ))
                            has_env_warn = True
                        elif code and not name:
                            name = f"Ambiente {code}"
                            warnings.append(ImportIssue(
                                sheet=sheet_ia,
                                row=row_num,
                                entity="environment",
                                severity="warning",
                                message=f"Ambiente sin nombre. Se usó 'Ambiente {code}'",
                                raw_value=code
                            ))
                            has_env_warn = True
                            
                        if not code:
                            errors.append(ImportIssue(
                                sheet=sheet_ia,
                                row=row_num,
                                entity="environment",
                                severity="error",
                                message="El número o nombre del ambiente es obligatorio para crearlo."
                            ))
                            has_env_err = True
                            
                        if not has_env_err:

                            environment_dict = {
                                "code": code[:50],
                                "name": name[:200],
                                "capacity": 0,
                                "environment_type": "fisico"
                            }
                            # Avoid duplicates in list
                            if not any(e["code"] == environment_dict["code"] for e in environments_list):
                                environments_list.append(environment_dict)
                                if has_env_warn:
                                    counts["environments"]["warnings"] += 1
                                else:
                                    counts["environments"]["valid"] += 1
                        else:
                            counts["environments"]["rejected"] += 1

                    # Parse Activity / RAP from row
                    actividad = row[actividad_idx] if actividad_idx is not None else None
                    if actividad and str(actividad).strip():
                        act_str = str(actividad).strip()
                        act_code = ""
                        act_desc = act_str
                        
                        match = re.match(r"^(\d+)\s*-\s*(.*)$", act_str)
                        if match:
                            act_code = match.group(1)
                            act_desc = match.group(2).strip()
                        else:
                            act_code = f"RAP-{get_stable_hash(act_str)}"
                            
                        try:
                            act_h = Decimal(str(row[horas_actividad_idx])) if (horas_actividad_idx is not None and row[horas_actividad_idx] is not None) else Decimal("0")
                        except Exception:
                            act_h = Decimal("0")
                            
                        rap_dict = {
                            "code": act_code[:50],
                            "description": act_desc[:500],
                            "estimated_hours": act_h,
                            "result_type": "lista_ia"
                        }
                        if not any(r["code"] == rap_dict["code"] for r in learning_results_list):
                            learning_results_list.append(rap_dict)
                            counts["learning_results"]["valid"] += 1

    # 2. Parse FICHAS
    sheet_fi = "FICHAS"
    if import_type in ("semaforos_sena", "groups"):
        if sheet_fi not in sheets_detected:
            errors.append(ImportIssue(
                sheet=sheet_fi,
                severity="error",
                message=f"La hoja obligatoria '{sheet_fi}' no fue encontrada en el archivo."
            ))
        else:
            sheet = wb[sheet_fi]
            rows = list(sheet.iter_rows(values_only=True))
            if len(rows) < 2:
                errors.append(ImportIssue(
                    sheet=sheet_fi,
                    severity="error",
                    message=f"La hoja '{sheet_fi}' está vacía o no contiene filas de datos."
                ))
            else:
                headers = [normalize_header(c) for c in rows[0]]
                
                # find indices
                code_idx = find_col_idx(headers, 'no_fichas', 'no_ficha', 'fichas')
                name_idx = find_col_idx(headers, 'ficha')
                level_idx = find_col_idx(headers, 'nivel')
                coordinacion_idx = find_col_idx(headers, 'coordinacion')
                tri_idx = find_col_idx(headers, 'tri', 'trimestre')
                start_date_idx = find_col_idx(headers, 'fecha_inicio', 'inicio')
                end_date_idx = find_col_idx(headers, 'fecha_fin_lectiva', 'fin_lectiva', 'fin')
                learners_idx = find_col_idx(headers, 'no_aprendices', 'aprendices')
                sede_idx = find_col_idx(headers, 'sede')
                jornada_idx = find_col_idx(headers, 'jornada')
                
                # Walk rows
                for r_idx in range(1, len(rows)):
                    row = rows[r_idx]
                    row_num = r_idx + 1
                    
                    if not any(c is not None for c in row):
                        continue
                        
                    has_group_err = False
                    has_group_warn = False
                    
                    code_val = row[code_idx] if code_idx is not None else None
                    name_val = row[name_idx] if name_idx is not None else None
                    level_val = row[level_idx] if level_idx is not None else None
                    coordinacion_val = row[coordinacion_idx] if coordinacion_idx is not None else ""
                    start_date_val = row[start_date_idx] if start_date_idx is not None else None
                    end_date_val = row[end_date_idx] if end_date_idx is not None else None
                    learners_val = row[learners_idx] if learners_idx is not None else None
                    sede_val = row[sede_idx] if sede_idx is not None else ""
                    jornada_val = row[jornada_idx] if jornada_idx is not None else ""
                    
                    if code_val is None or not str(code_val).strip():
                        errors.append(ImportIssue(
                            sheet=sheet_fi,
                            row=row_num,
                            entity="group",
                            severity="error",
                            message="El código de ficha (No. FICHAS) es obligatorio."
                        ))
                        has_group_err = True
                        
                    if not has_group_err:
                        code = str(code_val).strip()
                        name = str(name_val).strip() if name_val is not None else ""
                        
                        if not name:
                            name = f"Ficha {code}"
                            warnings.append(ImportIssue(
                                sheet=sheet_fi,
                                row=row_num,
                                entity="group",
                                severity="warning",
                                message=f"Ficha sin nombre. Se usó 'Ficha {code}'",
                                raw_value=code
                            ))
                            has_group_warn = True
                            
                        # learners count
                        try:
                            learners = int(str(learners_val)) if learners_val is not None else 0
                        except Exception:
                            learners = 0
                            warnings.append(ImportIssue(
                                sheet=sheet_fi,
                                row=row_num,
                                entity="group",
                                severity="warning",
                                message="Ficha sin número de aprendices. Se usó 0.",
                                raw_value=str(learners_val)
                            ))
                            has_group_warn = True
                            
                        # Dates
                        start_date = parse_excel_date(start_date_val)
                        end_date = parse_excel_date(end_date_val)
                        
                        # Program name extraction
                        prog_name = name
                        if "_" in name:
                            parts = name.split("_")
                            if len(parts) > 1:
                                prog_name = parts[-1].strip()
                                
                        prog_code = f"PROG-{get_stable_hash(prog_name)}"
                        prog_level = str(level_val).strip() if level_val is not None else "TECNOLOGO"
                        
                        program_dict = {
                            "code": prog_code,
                            "name": prog_name[:300],
                            "level": prog_level[:100]
                        }
                        if not any(p["code"] == program_dict["code"] for p in programs_list):
                            programs_list.append(program_dict)
                            
                        group_notes = []
                        if coordinacion_val:
                            group_notes.append(f"Coordinación: {coordinacion_val}")
                        if sede_val:
                            group_notes.append(f"Sede: {sede_val}")
                            
                        group_dict = {
                            "code": code[:50],
                            "name": name[:300],
                            "jornada": str(jornada_val).strip()[:50] if jornada_val else None,
                            "start_date": start_date,
                            "end_date": end_date,
                            "learners_count": learners,
                            "notes": " | ".join(group_notes) if group_notes else None,
                            "training_program_code": prog_code
                        }
                        
                        if not any(g["code"] == group_dict["code"] for g in groups_list):
                            groups_list.append(group_dict)
                            if has_group_warn:
                                counts["groups"]["warnings"] += 1
                            else:
                                counts["groups"]["valid"] += 1
                    else:
                        counts["groups"]["rejected"] += 1

    # 3. Parse Semáforos if present
    semaforo_sheets = ['Semaforo con RA', 'Semaforo Cadena', 'Semaforo RA - Oferta Abierta', 'Semaforo Oferta Abierta']
    for s_name in semaforo_sheets:
        if import_type != "semaforos_relacional" and s_name in sheets_detected:
            sheet = wb[s_name]
            rap_candidates = parse_semaforo_sheet(sheet)
            for rap in rap_candidates:
                desc = rap["description"]
                hours = rap["hours"]
                
                rap_code = ""
                match = re.match(r"^(\d+)\s*-\s*(.*)$", desc)
                if match:
                    rap_code = f"RAP-{match.group(1)}"
                    desc = match.group(2).strip()
                else:
                    rap_code = f"RAP-{get_stable_hash(desc)}"
                    
                rap_dict = {
                    "code": rap_code[:50],
                    "description": desc[:500],
                    "estimated_hours": Decimal(str(hours)) if hours is not None else None,
                    "result_type": "semaforo"
                }
                if not any(r["code"] == rap_dict["code"] for r in learning_results_list):
                    learning_results_list.append(rap_dict)
                    counts["learning_results"]["valid"] += 1

    if import_type == "semaforos_relacional":
        normalized_sheets = {"learning_results", "topics", "color_groups", "ra_topic_relations"}
        relational = (
            process_normalized_relational_workbook(wb)
            if normalized_sheets.issubset({normalize_header(name) for name in sheets_detected})
            else process_relational_semaforos(wb, warnings)
        )
        learning_results_list.extend(relational["learning_results"])
        topics_list = relational["topics"]
        color_groups_list = relational["color_groups"]
        ra_topic_relations_list = relational["ra_topic_relations"]
        counts["learning_results"]["valid"] += len(relational["learning_results"])
        counts["topics"]["valid"] = len(topics_list)
        counts["color_groups"]["valid"] = len(color_groups_list)
        counts["ra_topic_relations"]["valid"] = len(ra_topic_relations_list)

    instructors_list = deduplicate_records(
        instructors_list,
        lambda x: x.get("document_number", ""),
        "instructors",
        "LISTA INSTRUCTORES",
        warnings,
        errors,
    )
    environments_list = deduplicate_records(
        environments_list,
        lambda x: x.get("code", ""),
        "environments",
        "AMBIENTES",
        warnings,
        errors,
    )
    groups_list = deduplicate_records(
        groups_list,
        lambda x: x.get("code", ""),
        "groups",
        "FICHAS",
        warnings,
        errors,
    )
    programs_list = deduplicate_records(
        programs_list,
        lambda x: x.get("code", ""),
        "programs",
        "FICHAS",
        warnings,
        errors,
    )

    counts["instructors"]["valid"] = len(instructors_list)
    counts["environments"]["valid"] = len(environments_list)
    counts["groups"]["valid"] = len(groups_list)
    counts["programs"]["valid"] = len(programs_list)

    return {
        "import_type": import_type,
        "filename": filename,
        "sheets_detected": sheets_detected,
        "summary": {
            k: ImportEntitySummary(**v) for k, v in counts.items()
        },
        "items": {
            "academic_periods": [],
            "contract_types": [],
            "instructors": instructors_list,
            "environments": environments_list,
            "groups": groups_list,
            "programs": programs_list,
            "learning_results": learning_results_list,
            "topics": topics_list,
            "color_groups": color_groups_list,
            "ra_topic_relations": ra_topic_relations_list
        },
        "warnings": warnings,
        "errors": errors
    }


def _record_classification(
    entity_key: str,
    record: ClassifiedRecord,
    classified_counts: dict[str, dict[str, int]],
    conflicts_detail: list[ImportConflictItem],
    changes_detail: list[ImportChangeItem],
):
    if entity_key not in classified_counts:
        classified_counts[entity_key] = {"created": 0, "updated": 0, "unchanged": 0, "conflicts": 0}

    action_lower = record.action.lower()
    if action_lower == "create":
        classified_counts[entity_key]["created"] += 1
    elif action_lower == "update":
        classified_counts[entity_key]["updated"] += 1
        if record.changes:
            changes_detail.append(
                ImportChangeItem(
                    entity=entity_key,
                    natural_key=record.natural_key,
                    changes={
                        k: ImportFieldChange(before=v["before"], after=v["after"])
                        for k, v in record.changes.items()
                    },
                )
            )
    elif action_lower == "unchanged":
        classified_counts[entity_key]["unchanged"] += 1
    elif action_lower == "conflict":
        classified_counts[entity_key]["conflicts"] += 1
        conflicts_detail.append(
            ImportConflictItem(
                entity=entity_key,
                natural_key=record.natural_key,
                sheet=record.source_sheet,
                row=record.source_row,
                message=record.conflict_reason or "Conflicto con datos existentes",
            )
        )


def preview_import_for_coordination(
    session: Session,
    file_bytes: bytes,
    filename: str,
    import_type: str,
    coordination_id: Optional[int],
    is_global: bool = False,
) -> ImportPreviewResponse:
    res = process_workbook(file_bytes, filename, import_type)
    file_sha256 = compute_file_sha256(file_bytes)

    # Check if reimport
    is_reimport = False
    if coordination_id is not None:
        prev_batch = session.exec(
            select(ImportBatch).where(
                ImportBatch.coordination_id == coordination_id,
                ImportBatch.file_sha256 == file_sha256,
                ImportBatch.import_type == import_type,
                ImportBatch.status.in_(["completed", "completed_with_warnings"]),
            )
        ).first()
        if prev_batch:
            is_reimport = True
            res["warnings"].append(
                ImportIssue(
                    sheet="GENERAL",
                    severity="warning",
                    message="Este mismo archivo ya fue importado anteriormente para esta coordinación.",
                )
            )

    conflicts_detail: list[ImportConflictItem] = []
    changes_detail: list[ImportChangeItem] = []

    classified_counts: dict[str, dict[str, int]] = {
        k: {"created": 0, "updated": 0, "unchanged": 0, "conflicts": 0}
        for k in res["summary"].keys()
    }

    # 1. Academic Periods
    for item in res["items"].get("academic_periods", []):
        c = classify_academic_period(session, item, is_global)
        _record_classification("academic_periods", c, classified_counts, conflicts_detail, changes_detail)

    # 2. Programs
    for item in res["items"].get("programs", []):
        c = classify_training_program(session, item)
        _record_classification("programs", c, classified_counts, conflicts_detail, changes_detail)

    # 3. Contract types
    for item in res["items"].get("contract_types", []):
        c = classify_contract_type(session, item)
        _record_classification("contract_types", c, classified_counts, conflicts_detail, changes_detail)

    # 4. Instructors
    for item in res["items"].get("instructors", []):
        c = classify_instructor(session, item, coordination_id, is_global)
        _record_classification("instructors", c, classified_counts, conflicts_detail, changes_detail)

    # 5. Environments
    for item in res["items"].get("environments", []):
        c = classify_environment(session, item, coordination_id, is_global)
        _record_classification("environments", c, classified_counts, conflicts_detail, changes_detail)

    # 6. Groups
    for item in res["items"].get("groups", []):
        c = classify_group(session, item, coordination_id, is_global)
        _record_classification("groups", c, classified_counts, conflicts_detail, changes_detail)

    # 7. Learning Results, Topics, etc.
    for item in res["items"].get("learning_results", []):
        c = classify_generic_global_entity(session, LearningResult, "code", "learning_results", item)
        _record_classification("learning_results", c, classified_counts, conflicts_detail, changes_detail)

    for item in res["items"].get("topics", []):
        c = classify_generic_global_entity(session, Topic, "code", "topics", item)
        _record_classification("topics", c, classified_counts, conflicts_detail, changes_detail)

    for item in res["items"].get("ra_topic_relations", []):
        c = classify_generic_global_entity(session, LearningResultTopic, "relation_id", "ra_topic_relations", item)
        _record_classification("ra_topic_relations", c, classified_counts, conflicts_detail, changes_detail)

    # Update summaries
    for entity_name, counts in classified_counts.items():
        if entity_name in res["summary"]:
            curr_summary = res["summary"][entity_name]
            total_valid = counts["created"] + counts["updated"] + counts["unchanged"]
            curr_summary.valid = total_valid
            curr_summary.created = counts["created"]
            curr_summary.updated = counts["updated"]
            curr_summary.unchanged = counts["unchanged"]
            curr_summary.conflicts = counts["conflicts"]
            curr_summary.rejected += counts["conflicts"]

    return ImportPreviewResponse(
        import_type=res["import_type"],
        filename=res["filename"],
        sheets_detected=res["sheets_detected"],
        summary=res["summary"],
        items=res["items"],
        warnings=res["warnings"],
        errors=res["errors"],
        file_sha256=file_sha256,
        coordination_id=coordination_id,
        mode="safe_merge",
        is_reimport=is_reimport,
        conflicts_detail=conflicts_detail,
        changes_detail=changes_detail,
    )


def preview_workbook(file_bytes: bytes, filename: str, import_type: str) -> ImportPreviewResponse:
    res = process_workbook(file_bytes, filename, import_type)
    return ImportPreviewResponse(**res)


def commit_workbook_for_coordination(
    session: Session,
    file_bytes: bytes,
    import_type: str,
    coordination_id: Optional[int],
    user_id: int,
    filename: str = "import.xlsx",
    mode: str = "safe_merge",
    is_global: bool = False,
) -> ImportCommitResponse:
    file_sha256 = compute_file_sha256(file_bytes)
    res = process_workbook(file_bytes, filename, import_type)

    errors = res["errors"]
    if any(e.severity == "error" for e in errors):
        return ImportCommitResponse(
            status="failed",
            created={},
            updated={},
            unchanged={},
            conflicts={},
            rejected=len([e for e in errors if e.severity == "error"]),
            warnings=res["warnings"],
            errors=errors,
            file_sha256=file_sha256,
        )

    created_counts = {k: 0 for k in [
        "academic_periods", "contract_types", "instructors", "environments",
        "groups", "programs", "competencies", "learning_results", "topics",
        "learning_result_topics"
    ]}
    updated_counts = {k: 0 for k in created_counts.keys()}
    unchanged_counts = {k: 0 for k in created_counts.keys()}
    conflict_counts = {k: 0 for k in created_counts.keys()}

    try:
        # Create ImportBatch in the transaction
        batch = ImportBatch(
            import_type=import_type,
            coordination_id=coordination_id,
            uploaded_by_user_id=user_id,
            filename=filename,
            file_sha256=file_sha256,
            mode=mode,
            status="processing",
            created_at=datetime.utcnow(),
        )
        session.add(batch)
        session.flush()

        # 0. Academic Periods
        for period in res["items"].get("academic_periods", []):
            classified = classify_academic_period(session, period, is_global)
            if classified.action == "CONFLICT":
                conflict_counts["academic_periods"] += 1
                session.add(ImportBatchRecord(
                    batch_id=batch.id,
                    coordination_id=coordination_id,
                    entity_type="academic_periods",
                    natural_key=classified.natural_key,
                    action="CONFLICT",
                    source_sheet=classified.source_sheet,
                    source_row=classified.source_row,
                    conflict_reason=classified.conflict_reason,
                ))
                continue

            if classified.action == "CREATE" and is_global:
                payload = {
                    "year": period["year"],
                    "quarter_number": period["quarter_number"],
                    "name": period["name"],
                    "start_date": period["start_date"],
                    "end_date": period["end_date"],
                    "is_active": True,
                }
                obj = AcademicPeriod(**payload)
                session.add(obj)
                session.flush()
                created_counts["academic_periods"] += 1
                session.add(ImportBatchRecord(
                    batch_id=batch.id,
                    coordination_id=coordination_id,
                    entity_type="academic_periods",
                    entity_id=obj.id,
                    natural_key=classified.natural_key,
                    action="CREATE",
                    after_hash=compute_record_fingerprint(obj.model_dump()),
                    source_sheet=classified.source_sheet,
                    source_row=classified.source_row,
                ))
            elif classified.action == "UPDATE" and is_global:
                existing = classified.existing_obj
                for k, v in period.items():
                    merged = merge_field(getattr(existing, k, None), v)
                    setattr(existing, k, merged)
                session.add(existing)
                session.flush()
                updated_counts["academic_periods"] += 1
                session.add(ImportBatchRecord(
                    batch_id=batch.id,
                    coordination_id=coordination_id,
                    entity_type="academic_periods",
                    entity_id=existing.id,
                    natural_key=classified.natural_key,
                    action="UPDATE",
                    before_hash=classified.before_hash,
                    after_hash=compute_record_fingerprint(existing.model_dump()),
                    source_sheet=classified.source_sheet,
                    source_row=classified.source_row,
                ))
            else:
                unchanged_counts["academic_periods"] += 1

        # 1. Programs
        program_id_map = {}
        for prog in res["items"]["programs"]:
            classified = classify_training_program(session, prog)
            if classified.action == "CONFLICT":
                conflict_counts["programs"] += 1
                session.add(ImportBatchRecord(
                    batch_id=batch.id,
                    coordination_id=coordination_id,
                    entity_type="programs",
                    natural_key=classified.natural_key,
                    action="CONFLICT",
                    source_sheet=classified.source_sheet,
                    source_row=classified.source_row,
                    conflict_reason=classified.conflict_reason,
                ))
                continue

            payload = {
                "code": prog["code"],
                "name": prog["name"],
                "level": prog["level"],
                "is_active": True,
            }
            obj, created = upsert_entity(session, TrainingProgram, "code", prog["code"], payload)
            program_id_map[prog["code"]] = obj.id
            if created:
                created_counts["programs"] += 1
                act = "CREATE"
            else:
                act = "UNCHANGED"
                unchanged_counts["programs"] += 1
            session.add(ImportBatchRecord(
                batch_id=batch.id,
                coordination_id=coordination_id,
                entity_type="programs",
                entity_id=obj.id,
                natural_key=prog["code"],
                action=act,
                source_sheet="FICHAS",
                source_row=prog.get("source_row"),
            ))

        # 2. Groups (Fichas)
        for grp in res["items"]["groups"]:
            classified = classify_group(session, grp, coordination_id, is_global)
            if classified.action == "CONFLICT":
                conflict_counts["groups"] += 1
                session.add(ImportBatchRecord(
                    batch_id=batch.id,
                    coordination_id=coordination_id,
                    entity_type="groups",
                    natural_key=classified.natural_key,
                    action="CONFLICT",
                    source_sheet=classified.source_sheet,
                    source_row=classified.source_row,
                    conflict_reason=classified.conflict_reason,
                ))
                continue

            prog_id = program_id_map.get(grp["training_program_code"])
            if not prog_id:
                prog = session.exec(select(TrainingProgram).where(TrainingProgram.code == grp["training_program_code"])).first()
                prog_id = prog.id if prog else None

            if classified.action == "CREATE":
                payload = {
                    "code": grp["code"],
                    "name": grp["name"],
                    "jornada": grp["jornada"],
                    "modality": grp.get("modality"),
                    "trimester": grp.get("trimester"),
                    "start_date": grp["start_date"],
                    "end_date": grp["end_date"],
                    "productive_stage_start_date": grp.get("productive_stage_start_date"),
                    "productive_stage_end_date": grp.get("productive_stage_end_date"),
                    "learners_count": grp["learners_count"],
                    "notes": grp.get("notes"),
                    "training_program_id": prog_id,
                    "coordination_id": coordination_id,
                    "is_active": True,
                }
                new_group = Group(**payload)
                session.add(new_group)
                session.flush()
                created_counts["groups"] += 1
                session.add(ImportBatchRecord(
                    batch_id=batch.id,
                    coordination_id=coordination_id,
                    entity_type="groups",
                    entity_id=new_group.id,
                    natural_key=grp["code"],
                    action="CREATE",
                    after_hash=compute_record_fingerprint(new_group.model_dump()),
                    source_sheet=classified.source_sheet,
                    source_row=classified.source_row,
                ))
            elif classified.action == "UPDATE":
                existing = classified.existing_obj
                # Concurrency optimistic check
                curr_db_hash = compute_record_fingerprint(existing.model_dump())
                if classified.before_hash and curr_db_hash != classified.before_hash:
                    conflict_counts["groups"] += 1
                    session.add(ImportBatchRecord(
                        batch_id=batch.id,
                        coordination_id=coordination_id,
                        entity_type="groups",
                        entity_id=existing.id,
                        natural_key=grp["code"],
                        action="CONFLICT",
                        conflict_reason="CONFLICT_CONCURRENT_CHANGE: El registro fue modificado concurrentemente.",
                        source_sheet=classified.source_sheet,
                        source_row=classified.source_row,
                    ))
                    continue

                for fld, diff in classified.changes.items():
                    setattr(existing, fld, diff["after"])
                session.add(existing)
                session.flush()
                updated_counts["groups"] += 1
                session.add(ImportBatchRecord(
                    batch_id=batch.id,
                    coordination_id=coordination_id,
                    entity_type="groups",
                    entity_id=existing.id,
                    natural_key=grp["code"],
                    action="UPDATE",
                    before_hash=classified.before_hash,
                    after_hash=compute_record_fingerprint(existing.model_dump()),
                    changes=json.dumps(classified.changes, default=str),
                    source_sheet=classified.source_sheet,
                    source_row=classified.source_row,
                ))
            else:
                unchanged_counts["groups"] += 1

        # 3. Contract Types & Instructors
        contract_type_map = {}
        for item in res["items"].get("contract_types", []):
            payload = {
                "name": item["name"],
                "description": item.get("description"),
                "category": item.get("category"),
                "monthly_training_hours": item.get("monthly_training_hours", Decimal("0")),
                "monthly_additional_hours": item.get("monthly_additional_hours", Decimal("0")),
                "weekly_base_hours": item.get("weekly_base_hours", Decimal("0")),
                "weekly_max_hours": item.get("weekly_max_hours", Decimal("0")),
                "source_label": item.get("source_label"),
                "is_active": True,
            }
            obj, created = upsert_entity(session, ContractType, "name", item["name"], payload)
            contract_type_map[item["name"]] = obj.id
            if created:
                created_counts["contract_types"] += 1
            else:
                unchanged_counts["contract_types"] += 1

        for inst in res["items"]["instructors"]:
            classified = classify_instructor(session, inst, coordination_id, is_global)
            if classified.action == "CONFLICT":
                conflict_counts["instructors"] += 1
                session.add(ImportBatchRecord(
                    batch_id=batch.id,
                    coordination_id=coordination_id,
                    entity_type="instructors",
                    natural_key=classified.natural_key,
                    action="CONFLICT",
                    source_sheet=classified.source_sheet,
                    source_row=classified.source_row,
                    conflict_reason=classified.conflict_reason,
                ))
                continue

            ct_name = inst["contract_type_name"]
            if ct_name not in contract_type_map:
                ct_payload = {
                    "name": ct_name,
                    "description": f"Tipo de vinculacion {ct_name}",
                    "category": inst.get("contract_type_category"),
                    "monthly_training_hours": inst.get("monthly_training_hours", Decimal("0")),
                    "monthly_additional_hours": inst.get("monthly_additional_hours", Decimal("0")),
                    "weekly_base_hours": inst.get("contract_type_base_hours", Decimal("0")),
                    "weekly_max_hours": inst.get("contract_type_max_hours", Decimal("0")),
                    "source_label": ct_name,
                    "is_active": True,
                }
                ct_obj, ct_created = upsert_entity(session, ContractType, "name", ct_name, ct_payload)
                contract_type_map[ct_name] = ct_obj.id
                if ct_created:
                    created_counts["contract_types"] += 1

            ct_id = contract_type_map[ct_name]

            if classified.action == "CREATE":
                payload = {
                    "document_type": inst["document_type"],
                    "document_number": inst["document_number"],
                    "first_name": inst["first_name"],
                    "last_name": inst["last_name"],
                    "email": inst["email"],
                    "phone": inst.get("phone"),
                    "specialty": inst.get("specialty"),
                    "monthly_training_hours": inst.get("monthly_training_hours", Decimal("0")),
                    "monthly_additional_hours": inst.get("monthly_additional_hours", Decimal("0")),
                    "weekly_base_hours": inst["weekly_base_hours"],
                    "weekly_max_hours": inst["weekly_max_hours"],
                    "area": inst["area"],
                    "contract_type_id": ct_id,
                    "primary_coordination_id": coordination_id,
                    "is_active": True,
                }
                new_inst = Instructor(**payload)
                session.add(new_inst)
                session.flush()
                if coordination_id is not None:
                    session.add(InstructorCoordination(instructor_id=new_inst.id, coordination_id=coordination_id))
                    session.flush()
                created_counts["instructors"] += 1
                session.add(ImportBatchRecord(
                    batch_id=batch.id,
                    coordination_id=coordination_id,
                    entity_type="instructors",
                    entity_id=new_inst.id,
                    natural_key=inst["document_number"],
                    action="CREATE",
                    after_hash=compute_record_fingerprint(new_inst.model_dump()),
                    source_sheet=classified.source_sheet,
                    source_row=classified.source_row,
                ))
            else:
                existing = classified.existing_obj
                is_updated = False
                if classified.action == "UPDATE":
                    curr_db_hash = compute_record_fingerprint(existing.model_dump())
                    if classified.before_hash and curr_db_hash != classified.before_hash:
                        conflict_counts["instructors"] += 1
                        session.add(ImportBatchRecord(
                            batch_id=batch.id,
                            coordination_id=coordination_id,
                            entity_type="instructors",
                            entity_id=existing.id,
                            natural_key=inst["document_number"],
                            action="CONFLICT",
                            conflict_reason="CONFLICT_CONCURRENT_CHANGE",
                            source_sheet=classified.source_sheet,
                            source_row=classified.source_row,
                        ))
                        continue

                    for fld, diff in classified.changes.items():
                        setattr(existing, fld, diff["after"])
                    is_updated = bool(classified.changes)

                if coordination_id is not None:
                    link = session.exec(
                        select(InstructorCoordination).where(
                            InstructorCoordination.instructor_id == existing.id,
                            InstructorCoordination.coordination_id == coordination_id,
                        )
                    ).first()
                    if not link:
                        session.add(InstructorCoordination(instructor_id=existing.id, coordination_id=coordination_id))
                        is_updated = True

                session.add(existing)
                session.flush()
                if is_updated:
                    updated_counts["instructors"] += 1
                    act = "UPDATE"
                else:
                    unchanged_counts["instructors"] += 1
                    act = "UNCHANGED"

                session.add(ImportBatchRecord(
                    batch_id=batch.id,
                    coordination_id=coordination_id,
                    entity_type="instructors",
                    entity_id=existing.id,
                    natural_key=inst["document_number"],
                    action=act,
                    before_hash=classified.before_hash,
                    after_hash=compute_record_fingerprint(existing.model_dump()),
                    changes=json.dumps(classified.changes, default=str) if classified.changes else None,
                    source_sheet=classified.source_sheet,
                    source_row=classified.source_row,
                ))

        # 4. Environments
        for env in res["items"]["environments"]:
            classified = classify_environment(session, env, coordination_id, is_global)
            if classified.action == "CONFLICT":
                conflict_counts["environments"] += 1
                session.add(ImportBatchRecord(
                    batch_id=batch.id,
                    coordination_id=coordination_id,
                    entity_type="environments",
                    natural_key=classified.natural_key,
                    action="CONFLICT",
                    source_sheet=classified.source_sheet,
                    source_row=classified.source_row,
                    conflict_reason=classified.conflict_reason,
                ))
                continue

            if classified.action == "CREATE":
                payload = {
                    "code": env["code"],
                    "name": env["name"],
                    "location": env.get("location"),
                    "capacity": env.get("capacity", 0),
                    "environment_type": env.get("environment_type", "fisico"),
                    "notes": env.get("notes"),
                    "is_active": True,
                }
                new_env = Environment(**payload)
                session.add(new_env)
                session.flush()
                if coordination_id is not None:
                    session.add(EnvironmentCoordination(environment_id=new_env.id, coordination_id=coordination_id))
                    session.flush()
                created_counts["environments"] += 1
                session.add(ImportBatchRecord(
                    batch_id=batch.id,
                    coordination_id=coordination_id,
                    entity_type="environments",
                    entity_id=new_env.id,
                    natural_key=env["code"],
                    action="CREATE",
                    after_hash=compute_record_fingerprint(new_env.model_dump()),
                    source_sheet=classified.source_sheet,
                    source_row=classified.source_row,
                ))
            else:
                existing = classified.existing_obj
                is_updated = False
                if classified.action == "UPDATE":
                    curr_db_hash = compute_record_fingerprint(existing.model_dump())
                    if classified.before_hash and curr_db_hash != classified.before_hash:
                        conflict_counts["environments"] += 1
                        session.add(ImportBatchRecord(
                            batch_id=batch.id,
                            coordination_id=coordination_id,
                            entity_type="environments",
                            entity_id=existing.id,
                            natural_key=env["code"],
                            action="CONFLICT",
                            conflict_reason="CONFLICT_CONCURRENT_CHANGE",
                            source_sheet=classified.source_sheet,
                            source_row=classified.source_row,
                        ))
                        continue

                    for fld, diff in classified.changes.items():
                        setattr(existing, fld, diff["after"])
                    is_updated = bool(classified.changes)

                if coordination_id is not None:
                    link = session.exec(
                        select(EnvironmentCoordination).where(
                            EnvironmentCoordination.environment_id == existing.id,
                            EnvironmentCoordination.coordination_id == coordination_id,
                        )
                    ).first()
                    if not link:
                        session.add(EnvironmentCoordination(environment_id=existing.id, coordination_id=coordination_id))
                        is_updated = True

                session.add(existing)
                session.flush()
                if is_updated:
                    updated_counts["environments"] += 1
                    act = "UPDATE"
                else:
                    unchanged_counts["environments"] += 1
                    act = "UNCHANGED"

                session.add(ImportBatchRecord(
                    batch_id=batch.id,
                    coordination_id=coordination_id,
                    entity_type="environments",
                    entity_id=existing.id,
                    natural_key=env["code"],
                    action=act,
                    before_hash=classified.before_hash,
                    after_hash=compute_record_fingerprint(existing.model_dump()),
                    changes=json.dumps(classified.changes, default=str) if classified.changes else None,
                    source_sheet=classified.source_sheet,
                    source_row=classified.source_row,
                ))

        # 5. Generic Competencies and Learning Results (RAPs)
        raps_to_import = res["items"].get("learning_results", [])
        if raps_to_import and program_id_map and import_type not in ("semaforos_relacional", "schedule_normalized"):
            comp_map = {}
            for prog_code, prog_id in program_id_map.items():
                comp_code = f"COMP-GEN-{prog_code}"
                comp_payload = {
                    "code": comp_code,
                    "name": "Competencia Genérica para el programa",
                    "training_program_id": prog_id,
                    "is_active": True,
                }
                comp_obj, comp_created = upsert_entity(session, Competency, "code", comp_code, comp_payload)
                comp_map[prog_id] = comp_obj.id
                if comp_created:
                    created_counts["competencies"] += 1
                else:
                    unchanged_counts["competencies"] += 1

            for rap in raps_to_import:
                for prog_id, comp_id in comp_map.items():
                    rap_code = f"{rap['code']}-{prog_id}"
                    rap_payload = {
                        "code": rap_code[:50],
                        "description": rap["description"][:500],
                        "competency_id": comp_id,
                        "estimated_hours": rap["estimated_hours"],
                        "result_type": rap["result_type"],
                        "is_active": True,
                    }
                    obj, created = upsert_entity(session, LearningResult, "code", rap_code[:50], rap_payload)
                    if created:
                        created_counts["learning_results"] += 1
                    else:
                        unchanged_counts["learning_results"] += 1

        if import_type in ("semaforos_relacional", "schedule_normalized"):
            learning_result_id_map = {}
            for rap in raps_to_import:
                payload = {
                    "code": rap["code"][:50],
                    "description": rap["description"][:500],
                    "competency_id": None,
                    "estimated_hours": rap["estimated_hours"],
                    "result_type": rap["result_type"],
                    "is_active": True,
                }
                obj, created = upsert_entity(session, LearningResult, "code", rap["code"][:50], payload)
                learning_result_id_map[rap["code"]] = obj.id
                if created:
                    created_counts["learning_results"] += 1
                else:
                    unchanged_counts["learning_results"] += 1

            topic_id_map = {}
            for topic in res["items"].get("topics", []):
                payload = {
                    "code": topic["code"][:50],
                    "name": topic.get("name", topic.get("description", ""))[:500],
                    "program_scope": topic["program_scope"],
                    "trimester": topic["trimester"],
                    "trimester_number": topic["trimester_number"],
                    "estimated_hours": topic["estimated_hours"],
                    "source_sheet": topic["source_sheet"],
                    "source_address": topic["source_address"],
                    "source_row": topic["source_row"],
                    "source_col": topic["source_col"],
                    "color_key": topic["color_key"],
                    "color_hex": topic["color_hex"],
                    "is_active": True,
                }
                obj, created = upsert_entity(session, Topic, "code", topic["code"][:50], payload)
                topic_id_map[topic["code"]] = obj.id
                if created:
                    created_counts["topics"] += 1
                else:
                    unchanged_counts["topics"] += 1

            for relation in res["items"].get("ra_topic_relations", []):
                learning_result_id = learning_result_id_map.get(relation["learning_result_code"])
                topic_id = topic_id_map.get(relation["topic_code"])
                if learning_result_id is None or topic_id is None:
                    continue
                training_program_code = relation.get("training_program_code")
                training_program_id = program_id_map.get(training_program_code)
                if training_program_id is None and training_program_code:
                    program = session.exec(select(TrainingProgram).where(TrainingProgram.code == training_program_code)).first()
                    training_program_id = program.id if program else None
                payload = {
                    "relation_id": relation["relation_id"],
                    "training_program_id": training_program_id,
                    "training_program_code": training_program_code,
                    "training_program_name": relation.get("training_program_name"),
                    "learning_result_id": learning_result_id,
                    "topic_id": topic_id,
                    "group_id": relation["group_id"],
                    "program_scope": relation["program_scope"],
                    "trimester_number": relation["trimester_number"],
                    "color_key": relation["color_key"],
                    "color_hex": relation["color_hex"],
                    "relation_method": relation["relation_method"],
                    "relation_status": relation["relation_status"],
                    "confidence": relation["confidence"],
                    "needs_manual_review": relation["needs_manual_review"],
                }
                obj, created = upsert_entity(session, LearningResultTopic, "relation_id", relation["relation_id"], payload)
                if created:
                    created_counts["learning_result_topics"] += 1
                else:
                    unchanged_counts["learning_result_topics"] += 1

        total_created = sum(created_counts.values())
        total_updated = sum(updated_counts.values())
        total_unchanged = sum(unchanged_counts.values())
        total_conflicts = sum(conflict_counts.values())

        batch.created_count = total_created
        batch.updated_count = total_updated
        batch.unchanged_count = total_unchanged
        batch.conflict_count = total_conflicts
        batch.warning_count = len(res["warnings"])
        batch.error_count = len(res["errors"])
        batch.status = "completed_with_warnings" if (total_conflicts > 0 or res["warnings"]) else "completed"
        batch.committed_at = datetime.utcnow()
        batch.summary_data = json.dumps({
            "created": created_counts,
            "updated": updated_counts,
            "unchanged": unchanged_counts,
            "conflicts": conflict_counts,
        })
        session.flush()

        # SINGLE TRANSACTION COMMIT
        session.commit()

        return ImportCommitResponse(
            status=batch.status,
            batch_id=batch.id,
            created=created_counts,
            updated=updated_counts,
            unchanged=unchanged_counts,
            conflicts=conflict_counts,
            rejected=total_conflicts,
            warnings=res["warnings"],
            errors=errors,
            file_sha256=file_sha256,
        )

    except Exception:
        session.rollback()
        raise


def commit_workbook(session: Session, file_bytes: bytes, import_type: str, filename: str = "import.xlsx", mode: str = "upsert") -> ImportCommitResponse:
    admin_user = session.exec(select(User)).first()
    user_id = admin_user.id if admin_user else 1
    return commit_workbook_for_coordination(
        session=session,
        file_bytes=file_bytes,
        import_type=import_type,
        coordination_id=None,
        user_id=user_id,
        filename=filename,
        mode=mode,
        is_global=True,
    )

