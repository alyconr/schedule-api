from datetime import date, datetime, time
from decimal import Decimal
from io import BytesIO
from typing import Any
import unicodedata

from openpyxl import load_workbook
from sqlmodel import Session, select

from app.models import Environment, Group, Instructor, LearningResult, LearningResultTopic, Schedule, Topic
from app.schemas.imports import ImportCommitResponse, ImportEntitySummary, ImportIssue, ImportPreviewResponse
from app.schemas.schedules import ScheduleCreate
from app.services.schedule_period import quarter_for_date


def _header(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or "").strip().lower())
    return "_".join("".join(char for char in text if unicodedata.category(char) != "Mn").split())


def _date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value).strip()[:10])


def _time(value: Any, default: time | None = None) -> time:
    if value in (None, "") and default is not None:
        return default
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, (int, float)):
        minutes = round(float(value) * 24 * 60)
        return time((minutes // 60) % 24, minutes % 60)
    return time.fromisoformat(str(value).strip())


def _truthy(value: Any) -> bool:
    return _header(value) in {"1", "si", "sí", "true", "x"}


def _workbook_rows(file_bytes: bytes) -> tuple[str, list[tuple[int, dict[str, Any]]]]:
    workbook = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    sheet = workbook.active
    values = sheet.iter_rows(values_only=True)
    try:
        headers = [_header(value) for value in next(values)]
    except StopIteration:
        return sheet.title, []
    rows = [
        (row_number, dict(zip(headers, row, strict=False)))
        for row_number, row in enumerate(values, start=2)
        if any(value not in (None, "") for value in row)
    ]
    workbook.close()
    return sheet.title, rows


def _find_entities(session: Session, row: dict[str, Any]) -> tuple[Instructor | None, Group | None, Environment | None, LearningResult | None]:
    document = str(row.get("documento_instructor") or "").strip()
    group_code = str(row.get("ficha") or "").strip()
    environment_code = str(row.get("ambiente") or "").strip()
    learning_result_code = str(row.get("codigo_rap") or "").strip()
    instructor = session.exec(select(Instructor).where(Instructor.document_number == document)).first() if document else None
    group = session.exec(select(Group).where(Group.code == group_code)).first() if group_code else None
    environment = session.exec(select(Environment).where(Environment.code == environment_code)).first() if environment_code else None
    learning_result = session.exec(select(LearningResult).where(LearningResult.code == learning_result_code)).first() if learning_result_code else None
    return instructor, group, environment, learning_result


def _topic_assignment(
    session: Session,
    learning_result: LearningResult,
    raw_topic: Any,
) -> tuple[int | None, str | None]:
    topic_name = str(raw_topic or "").strip()
    relations = session.exec(
        select(LearningResultTopic, Topic)
        .join(Topic, Topic.id == LearningResultTopic.topic_id)
        .where(LearningResultTopic.learning_result_id == learning_result.id)
    ).all()
    if not relations:
        if not topic_name:
            raise ValueError("Debe indicar la temática del horario")
        return None, topic_name
    matches = [relation for relation, topic in relations if _header(topic.name) == _header(topic_name)]
    if len(matches) != 1:
        raise ValueError("La temática no coincide de forma única con el catálogo del RAP")
    return matches[0].id, None


def _parse_rows(
    session: Session,
    file_bytes: bytes,
    schedule_year: int,
    schedule_quarter: int,
) -> tuple[str, list[ScheduleCreate], list[ImportIssue]]:
    sheet, rows = _workbook_rows(file_bytes)
    parsed: list[ScheduleCreate] = []
    errors: list[ImportIssue] = []
    for row_number, row in rows:
        try:
            schedule_date = _date(row.get("fecha"))
            if schedule_date.year != schedule_year or quarter_for_date(schedule_date) != schedule_quarter:
                raise ValueError("La fecha no pertenece al año y trimestre seleccionados")
            instructor, group, environment, learning_result = _find_entities(session, row)
            if instructor is None:
                raise ValueError("No existe el instructor indicado en documento_instructor")
            is_additional = _truthy(row.get("es_hora_adicional"))
            if not is_additional and group is None:
                raise ValueError("No existe la ficha indicada")
            if not is_additional and environment is None:
                raise ValueError("No existe el ambiente indicado")
            if not is_additional and learning_result is None:
                raise ValueError("No existe el RAP indicado")
            topic_relation_id, manual_topic_name = (None, None)
            if not is_additional and learning_result is not None:
                topic_relation_id, manual_topic_name = _topic_assignment(
                    session,
                    learning_result,
                    row.get("tematica"),
                )
            start = _time(row.get("hora_inicio"), time(0, 0) if is_additional else None)
            end = _time(row.get("hora_fin"), time(0, 0) if is_additional else None)
            duration = row.get("duracion_horas")
            if duration in (None, ""):
                duration = ((end.hour * 60 + end.minute) - (start.hour * 60 + start.minute)) / 60
            payload = ScheduleCreate(
                instructor_id=instructor.id,
                group_id=None if is_additional else group.id,
                training_program_id=None if is_additional or group is None else group.training_program_id,
                competency_id=None if is_additional or learning_result is None else learning_result.competency_id,
                learning_result_id=None if is_additional else learning_result.id,
                learning_result_topic_id=topic_relation_id,
                manual_topic_name=manual_topic_name,
                environment_id=None if is_additional else environment.id,
                date=schedule_date,
                schedule_year=schedule_year,
                schedule_quarter=schedule_quarter,
                start_time=start,
                end_time=end,
                duration_hours=float(Decimal(str(duration))),
                is_additional_hours=is_additional,
                additional_hours_type=str(row.get("tipo_horas_adicionales") or "").strip() or None,
                notes=str(row.get("observaciones") or "").strip() or None,
            )
            parsed.append(payload)
        except Exception as exc:
            errors.append(
                ImportIssue(
                    sheet=sheet,
                    row=row_number,
                    entity="schedules",
                    severity="error",
                    message=str(exc),
                )
            )
    return sheet, parsed, errors


def preview_schedule_history(
    session: Session,
    file_bytes: bytes,
    filename: str,
    schedule_year: int,
    schedule_quarter: int,
) -> ImportPreviewResponse:
    sheet, parsed, errors = _parse_rows(session, file_bytes, schedule_year, schedule_quarter)
    return ImportPreviewResponse(
        import_type="schedule_history",
        filename=filename,
        sheets_detected=[sheet],
        summary={
            "schedules": ImportEntitySummary(valid=len(parsed), warnings=0, rejected=len(errors))
        },
        items={"schedules": [payload.model_dump(mode="json") for payload in parsed]},
        warnings=[],
        errors=errors,
    )


def commit_schedule_history(
    session: Session,
    file_bytes: bytes,
    filename: str,
    schedule_year: int,
    schedule_quarter: int,
) -> ImportCommitResponse:
    preview = preview_schedule_history(session, file_bytes, filename, schedule_year, schedule_quarter)
    if preview.errors:
        return ImportCommitResponse(
            status="failed",
            created={"schedules": 0},
            updated={},
            rejected=len(preview.errors),
            warnings=[],
            errors=preview.errors,
        )
    from app.api.routes.schedules import create_schedule

    created = 0
    errors: list[ImportIssue] = []
    for index, item in enumerate((preview.items or {}).get("schedules", []), start=2):
        try:
            result = create_schedule(ScheduleCreate.model_validate(item), session)
            if result.status == "blocked":
                raise ValueError("; ".join(validation.message for validation in result.validations))
            created += 1
        except Exception as exc:
            errors.append(
                ImportIssue(
                    sheet=preview.sheets_detected[0],
                    row=index,
                    entity="schedules",
                    severity="error",
                    message=str(exc),
                )
            )
    return ImportCommitResponse(
        status="completed" if not errors else "completed_with_errors",
        created={"schedules": created},
        updated={},
        rejected=len(errors),
        warnings=[],
        errors=errors,
    )
