import unittest
from datetime import date
from decimal import Decimal
from unittest.mock import MagicMock

from sqlmodel import Session, SQLModel, create_engine, select

from fastapi import HTTPException
from app.main import app
from app.models import (
    AcademicPeriod,
    ContractType,
    Environment,
    Group,
    Instructor,
    LearningResult,
    LearningResultTopic,
    Topic,
    TrainingProgram,
)
from app.services.import_service import (
    normalize_header,
    normalize_contract_type,
    vinculation_category,
    weekly_rule_hours_for_vinculation,
    parse_excel_date,
    get_stable_hash,
    build_program_code,
    program_code_from_name,
    find_col_idx,
    preview_workbook,
    commit_workbook
)
from app.services.imports import (
    canonical_program_key,
    parse_trimester_label,
)
from app.services.schedule_period import validate_schedule_period
from app.api.routes.imports import get_template_info
from app.api.routes.topics import _selection_statement, _to_selection_item


def build_schedule_normalized_workbook_bytes() -> bytes:
    import io
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LISTA INSTRUCTORES"
    ws.append(["NOMBRE COMPLETO", "TIPO DE VINCULACION", "HORAS FORMACION MES", "HORAS ADICIONALES", "COORDINACION"])
    ws.append(["ANA MARIA PEREZ", "CARRERA ADMINISTRATIVA", 170, 0, "Teleinformatica"])
    ws.append(["LUIS CARLOS RUIZ", "NOMBRAMIENTO ORDINARIO", 170, 4, "Teleinformatica"])
    ws.append(["MARTA SOFIA DIAZ", "NOMBRAMIENTO PROVISIONAL", 170, 0, "Teleinformatica"])
    ws.append(["JUAN PABLO GOMEZ", "CONTRATISTA SENA", 170, 0, "Teleinformatica"])

    ws = wb.create_sheet("AMBIENTES")
    ws.append(["NUMERO", "AMBIENTE O UBICACION"])
    ws.append(["301", "Aula 301"])

    ws = wb.create_sheet("FICHAS")
    ws.append([
        "No. FICHA",
        "FICHA",
        "NIVEL",
        "COORDINACION",
        "TRIMESTRE",
        "FECHA INICIO",
        "FECHA FIN LECTIVA",
        "FECHA INICIO PRODUCTIVA",
        "FECHA FIN PRODUCTIVA",
        "JORNADA",
        "SEDE",
    ])
    ws.append([
        "3068352",
        "DESARROLLO DE PROCESOS DE MERCADEO",
        "Tecnologo",
        "Mercadeo",
        "TRIMESTRE I - II",
        "2026-07-13",
        "2026-12-13",
        "2027-01-01",
        "2027-06-01",
        "Diurna",
        "Sede Colombia",
    ])

    headers = [
        "PROGRAMA DE FORMACION",
        "TRIMESTRE",
        "RESULTADO_APRENDIZAJE",
        "TIPO_RESULTADO_RA",
        "HORAS_SEMANA_RA",
        "HORAS_TRIMESTRE_RA",
        "TEMATICA",
        "TIPO_RESULTADO_TEMATICA",
        "HORAS_SEMANA_TEMATICA",
    ]
    for sheet_name, code, topic in (
        ("Semaforo con RA cadena", "RA1", "Investigacion de mercados"),
        ("Semaforo con RA Oferta Abierta", "RA2", "Segmentacion de clientes"),
    ):
        ws = wb.create_sheet(sheet_name)
        ws.append(headers)
        ws.append([
            "DESARROLLO DE PROCESOS DE MERCADEO",
            "TRIMESTRE I",
            f"Resultado {code}",
            "especifico",
            4,
            48,
            topic,
            "tematica",
            4,
        ])

    data = io.BytesIO()
    wb.save(data)
    return data.getvalue()


class ImportsRoutesAndServiceTest(unittest.TestCase):
    def _prefix_exists(self, method: str, prefix: str) -> bool:
        for route in app.routes:
            if hasattr(route, "methods") and hasattr(route, "path"):
                if route.path.startswith(prefix) and method in route.methods:
                    return True
        return False

    def test_routes_registered(self) -> None:
        self.assertTrue(self._prefix_exists("POST", "/api/v1/imports/preview"))
        self.assertTrue(self._prefix_exists("POST", "/api/v1/imports/commit"))
        self.assertTrue(self._prefix_exists("GET", "/api/v1/imports/template-info"))

    def test_normalize_header(self) -> None:
        self.assertEqual(normalize_header("INICIALES "), "iniciales")
        self.assertEqual(normalize_header("NOMBRE COMPLETO"), "nombre_completo")
        self.assertEqual(normalize_header("TIPO CONTRATO"), "tipo_contrato")
        self.assertEqual(normalize_header("ACTIVIDAD DE FORMACIÓN"), "actividad_de_formacion")
        self.assertEqual(normalize_header("No. FICHAS"), "no_fichas")
        self.assertEqual(normalize_header(None), "")

    def test_parse_excel_date_serial(self) -> None:
        # 45580 is 2024-10-15 in Excel (due to leap year bug logic)
        parsed_date = parse_excel_date(45580)
        self.assertEqual(parsed_date, date(2024, 10, 15))

    def test_parse_excel_date_string(self) -> None:
        self.assertEqual(parse_excel_date("2026-07-13"), date(2026, 7, 13))
        self.assertEqual(parse_excel_date("15/10/2024"), date(2024, 10, 15))

    def test_get_stable_hash(self) -> None:
        h1 = get_stable_hash("CRISTIAN BUITRAGO")
        h2 = get_stable_hash("CRISTIAN BUITRAGO")
        self.assertEqual(h1, h2)
        self.assertEqual(len(h1), 8)

    def test_program_code_ignores_accents_and_trailing_punctuation(self) -> None:
        self.assertEqual(
            program_code_from_name("ANÁLISIS Y DESARROLLO DE SOFTWARE."),
            program_code_from_name("ANALISIS Y DESARROLLO DE SOFTWARE"),
        )
        self.assertEqual(
            build_program_code("ANÁLISIS Y DESARROLLO DE SOFTWARE."),
            program_code_from_name("ANALISIS Y DESARROLLO DE SOFTWARE"),
        )

    def test_find_col_idx(self) -> None:
        headers = ["no_fichas", "ficha", "horasfromacion", "horas"]
        self.assertEqual(find_col_idx(headers, "horas"), 3)
        self.assertEqual(find_col_idx(headers, "ficha"), 1)
        self.assertEqual(find_col_idx(headers, "no_fichas"), 0)

    def test_normalize_contract_type_aliases(self) -> None:
        self.assertEqual(normalize_contract_type("CONTRATO")["name"], "contratista")
        self.assertEqual(normalize_contract_type("CONTRATISTA")["name"], "contratista")
        self.assertEqual(normalize_contract_type("CONTRATO SENA")["name"], "contratista")
        self.assertEqual(normalize_contract_type("PRESTACION DE SERVICIOS")["name"], "contratista")
        self.assertEqual(normalize_contract_type("CARRERA ADMINISTRATIVA")["name"], "planta")
        self.assertEqual(normalize_contract_type("NOMBRAMIENTO ORDINARIO")["name"], "planta")
        self.assertEqual(normalize_contract_type("NOMBRAMIENTO PROVISIONAL")["name"], "planta")
        self.assertEqual(normalize_contract_type("PLANTA")["name"], "planta")
        self.assertEqual(normalize_contract_type("SIN DATO")["name"], "otro")

    def test_vinculation_rules(self) -> None:
        self.assertEqual(vinculation_category("CARRERA ADMINISTRATIVA"), "planta")
        self.assertEqual(vinculation_category("NOMBRAMIENTO ORDINARIO"), "planta")
        self.assertEqual(vinculation_category("NOMBRAMIENTO PROVISIONAL"), "planta")
        self.assertEqual(vinculation_category("CONTRATISTA SENA"), "contratista")
        self.assertEqual(weekly_rule_hours_for_vinculation("CARRERA ADMINISTRATIVA"), Decimal("32"))
        self.assertEqual(weekly_rule_hours_for_vinculation("CONTRATISTA SENA"), Decimal("40"))

    def test_preview_does_not_truncate_items(self) -> None:
        import io
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "LISTA_INSTRUCTORES_AMBIENTES"
        ws.append(["NOMBRE COMPLETO", "INICIALES", "TIPO CONTRATO"])
        for i in range(55):
            ws.append([f"INSTRUCTOR {i}", f"I{i}", "CONTRATO"])

        f_bytes = io.BytesIO()
        wb.save(f_bytes)
        f_bytes.seek(0)

        res = preview_workbook(f_bytes.read(), "instructores.xlsx", "instructors_environments")
        self.assertEqual(len(res.items["instructors"]), 55)
        self.assertEqual(res.items["instructors"][0]["contract_type_name"], "contratista")

    def test_rejects_old_import_types(self) -> None:
        from app.api.routes.imports import ALLOWED_IMPORT_TYPES
        self.assertEqual(ALLOWED_IMPORT_TYPES, ["schedule_normalized", "schedule_history"])
        self.assertNotIn("semaforos_relacional", ALLOWED_IMPORT_TYPES)
        self.assertNotIn("semaforos_sena", ALLOWED_IMPORT_TYPES)

    def test_preview_schedule_normalized_workbook(self) -> None:
        res = preview_workbook(
            build_schedule_normalized_workbook_bytes(),
            "SEMAFOROS_NORMALIZADO_SCHEDULE_API.xlsx",
            "schedule_normalized",
        )

        self.assertEqual(res.errors, [])
        self.assertGreater(res.summary["instructors"].valid, 0)
        self.assertGreater(res.summary["environments"].valid, 0)
        self.assertGreater(res.summary["groups"].valid, 0)
        self.assertGreater(res.summary["programs"].valid, 0)
        self.assertEqual(res.summary["contract_types"].valid, 4)
        self.assertGreater(res.summary["learning_results"].valid, 0)
        self.assertGreater(res.summary["topics"].valid, 0)
        self.assertGreater(res.summary["ra_topic_relations"].valid, 0)
        instructor = res.items["instructors"][0]
        self.assertEqual(instructor["monthly_training_hours"], Decimal("170"))
        self.assertEqual(instructor["monthly_additional_hours"], Decimal("0"))
        self.assertEqual(instructor["weekly_base_hours"], Decimal("32"))
        self.assertEqual(instructor["weekly_max_hours"], Decimal("32"))
        contract_types = {item["name"]: item for item in res.items["contract_types"]}
        self.assertEqual(contract_types["CARRERA ADMINISTRATIVA"]["category"], "planta")
        self.assertEqual(contract_types["CARRERA ADMINISTRATIVA"]["weekly_base_hours"], Decimal("32"))
        self.assertEqual(contract_types["NOMBRAMIENTO ORDINARIO"]["weekly_max_hours"], Decimal("32"))
        self.assertEqual(contract_types["NOMBRAMIENTO PROVISIONAL"]["weekly_max_hours"], Decimal("32"))
        self.assertEqual(contract_types["CONTRATISTA SENA"]["category"], "contratista")
        self.assertEqual(contract_types["CONTRATISTA SENA"]["weekly_base_hours"], Decimal("40"))
        self.assertEqual(contract_types["CONTRATISTA SENA"]["monthly_training_hours"], Decimal("170"))
        self.assertEqual(res.items["programs"][0]["code"], build_program_code("DESARROLLO DE PROCESOS DE MERCADEO"))
        self.assertEqual(res.items["groups"][0]["trimester"], "TRIMESTRE I - II")
        self.assertNotIn("Trimestre:", res.items["groups"][0]["notes"] or "")
        self.assertIn("Sede: Sede Colombia", res.items["groups"][0]["notes"])
        relation = res.items["ra_topic_relations"][0]
        self.assertTrue(relation["learning_result_code"].startswith("CAD-TRIMESTRE_I-RAP-"))
        self.assertTrue(relation["topic_code"].startswith("TEM-CAD-"))
        self.assertTrue(relation["relation_id"].startswith("REL-CAD-"))
        self.assertIsNone(relation["color_key"])
        self.assertIsNone(relation["color_hex"])

    def test_commit_schedule_normalized_workbook(self) -> None:
        engine = create_engine("sqlite://", connect_args={"check_same_thread": False})
        SQLModel.metadata.create_all(engine)

        with Session(engine) as session:
            result = commit_workbook(
                session,
                build_schedule_normalized_workbook_bytes(),
                "schedule_normalized",
                filename="SEMAFOROS_NORMALIZADO_SCHEDULE_API.xlsx",
            )

            self.assertEqual(result.errors, [])
            contract_types = {ct.name: ct for ct in session.exec(select(ContractType)).all()}
            self.assertEqual(contract_types["CARRERA ADMINISTRATIVA"].category, "planta")
            self.assertEqual(contract_types["CARRERA ADMINISTRATIVA"].weekly_base_hours, Decimal("32.0"))
            self.assertEqual(contract_types["CONTRATISTA SENA"].category, "contratista")
            self.assertEqual(contract_types["CONTRATISTA SENA"].weekly_max_hours, Decimal("40.0"))
            self.assertEqual(contract_types["CONTRATISTA SENA"].monthly_training_hours, Decimal("170.0"))
            instructor = session.exec(select(Instructor).where(Instructor.document_number == f"TEMP-{get_stable_hash('ANA MARIA PEREZ')}")).first()
            self.assertIsNotNone(instructor)
            self.assertEqual(instructor.weekly_base_hours, Decimal("32.0"))
            self.assertEqual(instructor.monthly_training_hours, Decimal("170.0"))
            self.assertEqual(instructor.contract_type_id, contract_types["CARRERA ADMINISTRATIVA"].id)
            self.assertIsNotNone(session.exec(select(Environment)).first())
            self.assertIsNotNone(session.exec(select(TrainingProgram)).first())
            group = session.exec(select(Group)).first()
            self.assertIsNotNone(group)
            self.assertEqual(group.trimester, "TRIMESTRE I - II")
            self.assertNotIn("Trimestre:", group.notes or "")
            self.assertIn("Sede: Sede Colombia", group.notes)
            self.assertIsNotNone(group.productive_stage_start_date)
            self.assertIsNotNone(group.productive_stage_end_date)

            group.productive_stage_start_date = None
            group.productive_stage_end_date = None
            session.add(group)
            session.commit()
            second_result = commit_workbook(
                session,
                build_schedule_normalized_workbook_bytes(),
                "schedule_normalized",
                filename="SEMAFOROS_NORMALIZADO_SCHEDULE_API.xlsx",
            )
            session.refresh(group)
            self.assertGreater(second_result.updated["groups"], 0)
            self.assertEqual(group.trimester, "TRIMESTRE I - II")
            self.assertIsNotNone(group.productive_stage_start_date)
            self.assertIsNotNone(group.productive_stage_end_date)
            self.assertIsNotNone(session.exec(select(LearningResult)).first())
            self.assertIsNotNone(session.exec(select(Topic)).first())
            relation = session.exec(select(LearningResultTopic)).first()
            self.assertIsNotNone(relation)
            self.assertIsNotNone(relation.training_program_id)
            self.assertEqual(relation.training_program_code, session.exec(select(TrainingProgram)).first().code)
            self.assertEqual(relation.training_program_name, "DESARROLLO DE PROCESOS DE MERCADEO")
            self.assertIn(relation.program_scope, ("cadena", "oferta_abierta"))
            self.assertEqual(relation.relation_method, "normalized_excel_program_rap_topic_relation_without_color")
            self.assertEqual(relation.relation_status, "OK")
            self.assertEqual(relation.confidence, "alta")
            self.assertFalse(relation.needs_manual_review)
            row = session.exec(
                _selection_statement(
                    training_program_id=relation.training_program_id,
                    learning_result_id=relation.learning_result_id,
                )
            ).first()
            self.assertIsNotNone(row)
            topic_item = _to_selection_item(*row)
            self.assertEqual(topic_item.learning_result_topic_id, relation.id)
            self.assertEqual(topic_item.training_program_id, relation.training_program_id)

    def test_template_info_includes_only_schedule_normalized(self) -> None:
        info = get_template_info()
        self.assertEqual(info.supported_import_types, ["schedule_normalized", "schedule_history"])
        self.assertEqual(info.supported_formats, [".xlsx"])
        self.assertNotIn("semaforos_relacional", info.supported_import_types)
        self.assertEqual(info.required_sheets, ["LISTA INSTRUCTORES", "AMBIENTES", "FICHAS"])
        self.assertFalse(any(name.startswith("Semaforo con RA") for name in info.required_sheets))
        self.assertIn("TRIMESTRE", info.optional_sheets)
        self.assertIn("TEC CADENA", info.optional_sheets)
        self.assertIn("TEC REGULAR", info.optional_sheets)
        self.assertIn("TECNICO", info.optional_sheets)
        self.assertIn("AUXILIAR", info.optional_sheets)
        self.assertNotIn("LISTA_INSTRUCTORES_AMBIENTES", info.required_sheets)

    def test_parser_fails_when_mandatory_sheets_missing(self) -> None:
        import io
        import openpyxl
        
        # Create small empty workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        
        f_bytes = io.BytesIO()
        wb.save(f_bytes)
        f_bytes.seek(0)
        
        res = preview_workbook(f_bytes.read(), "dummy.xlsx", "semaforos_sena")
        self.assertTrue(len(res.errors) > 0)
    def test_import_row_missing_trimester(self) -> None:
        import io, openpyxl
        wb_bytes = build_schedule_normalized_workbook_bytes()
        wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
        ws = wb["FICHAS"]
        ws.cell(row=2, column=5, value="")
        buf = io.BytesIO()
        wb.save(buf)
        res = preview_workbook(buf.getvalue(), "test.xlsx", "schedule_normalized")
        self.assertTrue(len(res.errors) > 0)
        self.assertTrue(any(err.entity == "group" and "TRIMESTRE" in err.message for err in res.errors))
        self.assertNotIn("3068352", [g["code"] for g in res.items["groups"]])

    def test_canonical_program_key_and_prefix_stripping(self) -> None:
        self.assertEqual(
            canonical_program_key("TÉCNICO EN ASESORÍA COMERCIAL"),
            canonical_program_key("ASESORIA COMERCIAL"),
        )
        self.assertEqual(
            build_program_code("TÉCNICO EN ASESORÍA COMERCIAL"),
            build_program_code("ASESORIA COMERCIAL"),
        )
        self.assertEqual(
            build_program_code("TECNÓLOGO EN DESARROLLO DE SOFTWARE"),
            build_program_code("DESARROLLO DE SOFTWARE"),
        )
        self.assertEqual(
            build_program_code("AUXILIAR EN PROMOCION DE PRODUCTOS"),
            build_program_code("PROMOCION DE PRODUCTOS"),
        )

    def test_roman_trimesters_extended(self) -> None:
        self.assertEqual(parse_trimester_label("TRIMESTRE I")[1], 1)
        self.assertEqual(parse_trimester_label("TRIMESTRE VII")[1], 7)
        self.assertEqual(parse_trimester_label("VII TRIMESTRE")[1], 7)
        self.assertEqual(parse_trimester_label("TRIMESTRE XII")[1], 12)
        self.assertEqual(parse_trimester_label("7")[1], 7)
        self.assertEqual(parse_trimester_label("TRIMESTRE 7")[1], 7)
        self.assertEqual(parse_trimester_label("TRIMESTRE I - II")[1], 1)

    def test_preview_schedule_normalized_v2(self) -> None:
        content = build_schedule_normalized_v2_workbook_bytes()
        result = preview_workbook(
            content,
            "SEMAFOROS_NORMALIZADO_SCHEDULE_API.xlsx",
            "schedule_normalized",
        )

        self.assertEqual(result.errors, [])
        self.assertIn("FICHAS", result.sheets_detected)
        self.assertIn("LISTA INSTRUCTORES", result.sheets_detected)
        self.assertIn("AMBIENTES", result.sheets_detected)
        self.assertIn("TRIMESTRE", result.sheets_detected)
        self.assertIn("TEC CADENA", result.sheets_detected)
        self.assertIn("TEC REGULAR", result.sheets_detected)
        self.assertIn("TECNICO", result.sheets_detected)
        self.assertIn("AUXILIAR", result.sheets_detected)

        self.assertGreater(result.summary["instructors"].valid, 0)
        self.assertGreater(result.summary["environments"].valid, 0)
        self.assertGreater(result.summary["groups"].valid, 0)
        self.assertGreater(result.summary["programs"].valid, 0)
        self.assertGreater(result.summary["learning_results"].valid, 0)
        self.assertGreater(result.summary["topics"].valid, 0)
        self.assertGreater(result.summary["ra_topic_relations"].valid, 0)
        self.assertGreater(result.summary["academic_periods"].valid, 0)

        # Trimester VII resolution
        tri_vii_groups = [g for g in result.items["groups"] if g["code"] == "3068002"]
        self.assertEqual(len(tri_vii_groups), 1)
        self.assertEqual(tri_vii_groups[0]["trimester"], "TRIMESTRE VII")

        # Instructor with real document
        inst_real = next(i for i in result.items["instructors"] if i["document_number"] == "1017123456")
        self.assertEqual(inst_real["email"], "carlos.ramirez@sena.edu.co")
        self.assertEqual(inst_real["phone"], "3101234567")
        self.assertEqual(inst_real["specialty"], "Sistemas e Informatica")

        # Instructor without document gets TEMP-* and warning
        inst_temp = next(i for i in result.items["instructors"] if i["document_number"].startswith("TEMP-"))
        self.assertIsNotNone(inst_temp)
        self.assertTrue(any(w.entity == "instructor" and "documento temporal" in w.message for w in result.warnings))

        # Environment capacity and location
        env_301 = next(e for e in result.items["environments"] if e["code"] == "301")
        self.assertEqual(env_301["capacity"], 35)
        self.assertEqual(env_301["location"], "Sede Central")
        self.assertEqual(env_301["environment_type"], "fisico")

        env_virt = next(e for e in result.items["environments"] if e["code"] == "VIRT-1")
        self.assertEqual(env_virt["environment_type"], "virtual")

        # Group modality
        grp_virt = next(g for g in result.items["groups"] if g["code"] == "3068003")
        self.assertEqual(grp_virt["modality"], "Virtual")

        # Excel formula error #NAME? in AUXILIAR generated warning, not error
        formula_warnings = [w for w in result.warnings if w.sheet == "AUXILIAR" and "#NAME?" in (w.raw_value or "")]
        self.assertGreater(len(formula_warnings), 0)

        # Cross-program RAP collision test: same description in DESARROLLO DE SOFTWARE vs ASESORIA COMERCIAL
        ra_sw = next(
            r for r in result.items["learning_results"]
            if r["description"] == "Atencion al cliente y requerimientos" and "CAD" not in r["code"] and "TEC" not in r["code"]
        )
        ra_com = next(
            r for r in result.items["learning_results"]
            if r["description"] == "Atencion al cliente y requerimientos" and "TEC" in r["code"]
        )
        self.assertNotEqual(ra_sw["code"], ra_com["code"])

    def test_commit_schedule_normalized_v2(self) -> None:
        engine = create_engine("sqlite://", connect_args={"check_same_thread": False})
        SQLModel.metadata.create_all(engine)

        with Session(engine) as session:
            result = commit_workbook(
                session,
                build_schedule_normalized_v2_workbook_bytes(),
                "schedule_normalized",
                filename="SEMAFOROS_NORMALIZADO_SCHEDULE_API.xlsx",
            )

            self.assertEqual(result.errors, [])
            self.assertGreater(result.created["academic_periods"], 0)
            self.assertGreater(result.created["groups"], 0)
            self.assertGreater(result.created["instructors"], 0)
            self.assertGreater(result.created["environments"], 0)
            self.assertGreater(result.created["learning_results"], 0)
            self.assertGreater(result.created["topics"], 0)
            self.assertGreater(result.created["learning_result_topics"], 0)

            # Check AcademicPeriod in DB
            periods = session.exec(select(AcademicPeriod)).all()
            self.assertEqual(len(periods), 4)
            q1 = next(p for p in periods if p.quarter_number == 1)
            self.assertEqual(q1.start_date, date(2026, 1, 29))
            self.assertEqual(q1.end_date, date(2026, 4, 14))

            # Check Instructor in DB
            instructor = session.exec(select(Instructor).where(Instructor.document_number == "1017123456")).first()
            self.assertIsNotNone(instructor)
            self.assertEqual(instructor.phone, "3101234567")
            self.assertEqual(instructor.specialty, "Sistemas e Informatica")

            # Check Group in DB
            group = session.exec(select(Group).where(Group.code == "3068003")).first()
            self.assertIsNotNone(group)
            self.assertEqual(group.modality, "Virtual")

            # Check Environment in DB
            env = session.exec(select(Environment).where(Environment.code == "301")).first()
            self.assertIsNotNone(env)
            self.assertEqual(env.capacity, 35)

            # Check LearningResultTopic in DB
            relation = session.exec(select(LearningResultTopic)).first()
            self.assertIsNotNone(relation)
            self.assertIsNotNone(relation.training_program_id)

    def test_reimport_schedule_normalized_v2_is_idempotent(self) -> None:
        engine = create_engine("sqlite://", connect_args={"check_same_thread": False})
        SQLModel.metadata.create_all(engine)

        with Session(engine) as session:
            v2_bytes = build_schedule_normalized_v2_workbook_bytes()
            res1 = commit_workbook(session, v2_bytes, "schedule_normalized")
            self.assertEqual(res1.errors, [])

            total_periods_1 = len(session.exec(select(AcademicPeriod)).all())
            total_groups_1 = len(session.exec(select(Group)).all())
            total_ras_1 = len(session.exec(select(LearningResult)).all())
            total_topics_1 = len(session.exec(select(Topic)).all())
            total_rels_1 = len(session.exec(select(LearningResultTopic)).all())

            # Second import of exact same workbook
            res2 = commit_workbook(session, v2_bytes, "schedule_normalized")
            self.assertEqual(res2.errors, [])

            total_periods_2 = len(session.exec(select(AcademicPeriod)).all())
            total_groups_2 = len(session.exec(select(Group)).all())
            total_ras_2 = len(session.exec(select(LearningResult)).all())
            total_topics_2 = len(session.exec(select(Topic)).all())
            total_rels_2 = len(session.exec(select(LearningResultTopic)).all())

            self.assertEqual(total_periods_1, total_periods_2)
            self.assertEqual(total_groups_1, total_groups_2)
            self.assertEqual(total_ras_1, total_ras_2)
            self.assertEqual(total_topics_1, total_topics_2)
            self.assertEqual(total_rels_1, total_rels_2)
            self.assertGreater(res2.updated["groups"], 0)

    def test_institutional_period_validation_vs_calendar(self) -> None:
        engine = create_engine("sqlite://", connect_args={"check_same_thread": False})
        SQLModel.metadata.create_all(engine)

        with Session(engine) as session:
            # Commit v2 so academic periods exist
            commit_workbook(session, build_schedule_normalized_v2_workbook_bytes(), "schedule_normalized")

            # In institutional period: 2026-04-10 is inside Q1 (2026-01-29 to 2026-04-14)
            # Validating with session -> SUCCESS
            validate_schedule_period(date(2026, 4, 10), 2026, 1, session=session)

            # Validating without session (calendar fallback: month 4 is calendar Q2, not Q1) -> FAILS 422
            with self.assertRaises(HTTPException) as ctx:
                validate_schedule_period(date(2026, 4, 10), 2026, 1, session=None)
            self.assertEqual(ctx.exception.status_code, 422)
            self.assertIn("no pertenece al trimestre 1", ctx.exception.detail)

            # Date after period end (2026-04-16 is after 2026-04-14) with session -> FAILS 422
            with self.assertRaises(HTTPException) as ctx_inst:
                validate_schedule_period(date(2026, 4, 16), 2026, 1, session=session)
            self.assertEqual(ctx_inst.exception.status_code, 422)
            self.assertIn("no pertenece al periodo institucional", ctx_inst.exception.detail)


def build_schedule_normalized_v2_workbook_bytes() -> bytes:
    import io
    import openpyxl

    wb = openpyxl.Workbook()

    # 1. FICHAS
    ws = wb.active
    ws.title = "FICHAS"
    ws.append([
        "No. FICHA",
        "NOMBRE DEL PROGRAMA DE FORMACIÓN",
        "NIVEL",
        "COORDINACION",
        "TRIMESTRE",
        "INICIO LECTIVA",
        "FIN DE ETAPA LECTIVA",
        "INICIO PRODUCTIVA",
        "FIN PRODUCTIVA",
        "JORNADA",
        "SEDE",
        "TIPO DE RESPUESTA",
        "Modalidad",
    ])
    ws.append([
        "3068001",
        "DESARROLLO DE SOFTWARE",
        "Tecnologo",
        "Teleinformatica",
        "TRIMESTRE I",
        "2026-01-29",
        "2026-04-14",
        "2026-04-20",
        "2026-10-20",
        "Diurna",
        "Sede Central",
        "ABIERTA",
        "Presencial",
    ])
    ws.append([
        "3068002",
        "GESTION CONTABLE",
        "Tecnologo",
        "Finanzas",
        "TRIMESTRE VII",
        "2026-04-20",
        "2026-07-04",
        "2026-07-05",
        "2026-12-31",
        "Nocturna",
        "Sede Norte",
        "CERRADA",
        "Presencial",
    ])
    ws.append([
        "3068003",
        "TÉCNICO EN ASESORÍA COMERCIAL",
        "Tecnico",
        "Comercio",
        "TRIMESTRE I",
        "2026-01-29",
        "2026-04-14",
        "2026-04-20",
        "2026-07-20",
        "Diurna",
        "Sede Central",
        "ABIERTA",
        "Virtual",
    ])
    ws.append([
        "3068004",
        "AUXILIAR EN PROMOCION DE PRODUCTOS",
        "Auxiliar",
        "Mercadeo",
        "TRIMESTRE I",
        "2026-01-29",
        "2026-04-14",
        "2026-04-20",
        "2026-06-20",
        "Diurna",
        "Sede Sur",
        "CERRADA",
        "Presencial",
    ])

    # 2. LISTA INSTRUCTORES
    ws = wb.create_sheet("LISTA INSTRUCTORES")
    ws.append([
        "TIPO",
        "NUMERO",
        "NOMBRE COMPLETO",
        "TIPO DE VINCULACIÓN",
        "HORAS FORMACION MES",
        "HORAS ADICIONALES",
        "COORDINACION",
        "ESPECILIALIDAD O AREA",
        "CORREO",
        "TELEFONO",
    ])
    ws.append([
        "CC",
        "1017123456",
        "CARLOS ANDRES RAMIREZ",
        "CARRERA ADMINISTRATIVA",
        170,
        0,
        "Teleinformatica",
        "Sistemas e Informatica",
        "carlos.ramirez@sena.edu.co",
        "3101234567",
    ])
    ws.append([
        "",
        "",
        "MARIA FERNANDA LOPEZ",
        "CONTRATISTA SENA",
        160,
        0,
        "Teleinformatica",
        "Redes",
        "",
        "",
    ])

    # 3. AMBIENTES
    ws = wb.create_sheet("AMBIENTES")
    ws.append(["NUMERO", "SEDE", "CAPACIDAD", "COORDINACIÓN", "TIPO"])
    ws.append(["301", "Sede Central", 35, "Teleinformatica", "FISICO"])
    ws.append(["VIRT-1", "Sede Virtual", 50, "Teleinformatica", "VIRTUAL"])

    # 4. TRIMESTRE
    ws = wb.create_sheet("TRIMESTRE")
    ws.append(["Nombre", "Fecha Inicio", "Fecha Fin"])
    ws.append(["I Trimestre", "2026-01-29", "2026-04-14"])
    ws.append(["II Trimestre", "2026-04-20", "2026-07-04"])
    ws.append(["III Trimestre", "2026-07-09", "2026-09-28"])
    ws.append(["IV Trimestre", "2026-10-02", "2026-12-16"])

    sem_headers = [
        "PROGRAMA DE FORMACION",
        "TRIMESTRE",
        "RESULTADO_APRENDIZAJE",
        "TIPO_RESULTADO_RA (completo-parcial)",
        "HORAS_SEMANA_RA",
        "HORAS_TRIMESTRE_RA",
        "TEMATICA",
        "TIPO_RESULTADO_TEMATICA (completo-parcial)",
        "HORAS_SEMANA_TEMATICA",
    ]

    # 5. TEC CADENA
    ws = wb.create_sheet("TEC CADENA")
    ws.append(sem_headers)
    ws.append([
        "GESTION CONTABLE",
        "TRIMESTRE VII",
        "Auditoria de estados financieros",
        "completo",
        6,
        60,
        "Procedimientos de control interno",
        "completo",
        6,
    ])

    # 6. TEC REGULAR
    ws = wb.create_sheet("TEC REGULAR")
    ws.append(sem_headers)
    ws.append([
        "DESARROLLO DE SOFTWARE",
        "TRIMESTRE I",
        "Construir algoritmos estructurados",
        "completo",
        4,
        48,
        "Estructuras de control y datos",
        "completo",
        4,
    ])
    ws.append([
        "DESARROLLO DE SOFTWARE",
        "TRIMESTRE I",
        "Atencion al cliente y requerimientos",
        "parcial",
        4,
        48,
        "Levantamiento de requerimientos",
        "parcial",
        4,
    ])

    # 7. TECNICO
    ws = wb.create_sheet("TECNICO")
    ws.append(sem_headers)
    ws.append([
        "ASESORIA COMERCIAL",  # Canonical matching tests prefix stripping
        "TRIMESTRE I",
        "Atencion al cliente y requerimientos",  # Same text as in DESARROLLO DE SOFTWARE
        "completo",
        5,
        50,
        "Tecnicas de comunicacion asertiva",
        "completo",
        5,
    ])

    # 8. AUXILIAR
    ws = wb.create_sheet("AUXILIAR")
    ws.append(sem_headers)
    ws.append([
        "AUXILIAR EN PROMOCION DE PRODUCTOS",
        "TRIMESTRE I",
        "Organizacion de exhibiciones comerciales",
        "completo",
        8,
        "#NAME?",  # Formula error in hours cell
        "Merchandising visual en punto de venta",
        "completo",
        8,
    ])

    data = io.BytesIO()
    wb.save(data)
    return data.getvalue()


if __name__ == "__main__":
    unittest.main()

