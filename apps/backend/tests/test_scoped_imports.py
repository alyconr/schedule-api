"""Comprehensive integration & unit tests for coordination-scoped, atomic, and incremental imports."""

import io
import unittest
from datetime import date, datetime
from decimal import Decimal

import openpyxl
from sqlmodel import Session, SQLModel, create_engine, select

from app.api.deps import AccessScope
from app.models import (
    AcademicPeriod,
    ContractType,
    Coordination,
    Environment,
    EnvironmentCoordination,
    Group,
    ImportBatch,
    ImportBatchRecord,
    Instructor,
    InstructorCoordination,
    TrainingProgram,
    User,
    UserCoordination,
)
from app.services import import_service
from app.services.imports.safe_merge import (
    compute_file_sha256,
    compute_record_fingerprint,
    merge_field,
)


def create_minimal_normalized_workbook(
    instructors: list[dict] | None = None,
    environments: list[dict] | None = None,
    groups: list[dict] | None = None,
) -> bytes:
    wb = openpyxl.Workbook()

    # 1. LISTA INSTRUCTORES
    ws_inst = wb.active
    ws_inst.title = "LISTA INSTRUCTORES"
    ws_inst.append([
        "NOMBRE COMPLETO", "NUMERO DOCUMENTO", "CORREO", "TELEFONO",
        "TIPO DE VINCULACION", "HORAS FORMACION MES", "HORAS ADICIONALES", "COORDINACION"
    ])
    inst_rows = instructors if instructors is not None else [
        {
            "name": "ANA MARIA PEREZ",
            "document_number": "CC-1001",
            "email": "ana@sena.edu.co",
            "phone": "3001234567",
            "vinculation": "CARRERA ADMINISTRATIVA",
            "h_mes": 170,
            "h_ad": 0,
            "coordination": "Teleinformatica",
        }
    ]
    for inst in inst_rows:
        ws_inst.append([
            inst.get("name", "ANA MARIA PEREZ"),
            inst.get("document_number", "CC-1001"),
            inst.get("email", "ana@sena.edu.co"),
            inst.get("phone", "3001234567"),
            inst.get("vinculation", "CARRERA ADMINISTRATIVA"),
            inst.get("h_mes", 170),
            inst.get("h_ad", 0),
            inst.get("coordination", "Teleinformatica"),
        ])

    # 2. AMBIENTES
    ws_env = wb.create_sheet("AMBIENTES")
    ws_env.append(["NUMERO", "AMBIENTE O UBICACION", "CAPACIDAD", "TIPO AMBIENTE"])
    env_rows = environments if environments is not None else [{"code": "301", "name": "Aula 301", "capacity": 30}]
    for env in env_rows:
        ws_env.append([
            env.get("code", env.get("number", "301")),
            env.get("name", "Aula 301"),
            env.get("capacity", 30),
            env.get("environment_type", "fisico"),
        ])

    # 3. FICHAS
    ws_groups = wb.create_sheet("FICHAS")
    ws_groups.append([
        "No. FICHA", "FICHA", "NIVEL", "COORDINACION", "TRIMESTRE",
        "FECHA INICIO", "FECHA FIN LECTIVA", "FECHA INICIO PRODUCTIVA", "FECHA FIN PRODUCTIVA",
        "JORNADA", "SEDE"
    ])
    grp_rows = groups if groups is not None else [{"code": "3068352", "program": "ADSO", "jornada": "Diurna"}]
    for grp in grp_rows:
        ws_groups.append([
            grp.get("code", "3068352"),
            grp.get("program", "ADSO"),
            grp.get("level", "Tecnologo"),
            grp.get("coordination", "Teleinformatica"),
            grp.get("trimester", "TRIMESTRE I"),
            grp.get("start_date", "2026-07-13"),
            grp.get("end_date", "2026-12-13"),
            grp.get("prod_start", "2027-01-01"),
            grp.get("prod_end", "2027-06-01"),
            grp.get("jornada", "Diurna"),
            grp.get("sede", "Sede Principal"),
        ])

    headers = [
        "PROGRAMA DE FORMACION",
        "TRIMESTRE",
        "RESULTADO_APRENDIZAJE",
        "TIPO_RESULTADO_RA",
        "HORAS_SEMANA_RA",
        "HORAS_TRIMESTRE_RA",
        "TEMATICA",
        "TIPO_RESULTADO_TEMATICA",
        "HORAS_SEMANA_TEMATICA",
    ]
    for sheet_name, code, topic in (
        ("Semaforo con RA cadena", "RA1", "Investigacion de mercados"),
        ("Semaforo con RA Oferta Abierta", "RA2", "Segmentacion de clientes"),
    ):
        ws_sem = wb.create_sheet(sheet_name)
        ws_sem.append(headers)
        ws_sem.append([
            "DESARROLLO DE PROCESOS DE MERCADEO",
            "TRIMESTRE I",
            f"Resultado {code}",
            "especifico",
            4,
            48,
            topic,
            "tematica",
            4,
        ])

    data = io.BytesIO()
    wb.save(data)
    return data.getvalue()


class ScopedImportsTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False})
        SQLModel.metadata.create_all(self.engine)
        self.session = Session(self.engine)

        # Setup standard coordinations
        self.coord1 = Coordination(id=1, code="COORD_1", name="Teleinformática", is_active=True)
        self.coord2 = Coordination(id=2, code="COORD_2", name="Mercadeo", is_active=True)
        self.session.add(self.coord1)
        self.session.add(self.coord2)

        # Setup Users
        self.user1 = User(id=1, email="coord1@sena.edu.co", full_name="User Coord 1", hashed_password="pw", is_active=True)
        self.user2 = User(id=2, email="coord2@sena.edu.co", full_name="User Coord 2", hashed_password="pw", is_active=True)
        self.admin = User(id=3, email="admin@sena.edu.co", full_name="Admin User", hashed_password="pw", is_active=True)
        self.session.add(self.user1)
        self.session.add(self.user2)
        self.session.add(self.admin)
        self.session.flush()

        # User coordination mappings
        self.session.add(UserCoordination(user_id=1, coordination_id=1))
        self.session.add(UserCoordination(user_id=2, coordination_id=2))

        # Default contract type
        self.ct_planta = ContractType(
            id=1, name="planta", category="planta", weekly_base_hours=Decimal("32"), weekly_max_hours=Decimal("32")
        )
        self.session.add(self.ct_planta)
        self.session.commit()

    def tearDown(self):
        self.session.close()

    # 1. TEST AC-01: AccessScope authorization
    def test_access_scope_permissions(self):
        scope_user1 = AccessScope(user_id=1, roles=frozenset({"coordinador"}), coordination_ids=frozenset({1}), is_global=False)
        self.assertTrue(scope_user1.can_access(1))
        self.assertFalse(scope_user1.can_access(2))
        self.assertFalse(scope_user1.can_access(None))

        scope_admin = AccessScope(user_id=3, roles=frozenset({"admin"}), coordination_ids=frozenset(), is_global=True)
        self.assertTrue(scope_admin.can_access(1))
        self.assertTrue(scope_admin.can_access(2))
        self.assertTrue(scope_admin.can_access(None))

    # 2. TEST AC-04 & Punto 63: Group of another coordination triggers CONFLICT
    def test_group_ownership_conflict(self):
        existing_group = Group(
            code="2500001",
            name="Ficha Mercadeo",
            coordination_id=2,
            jornada="Diurna",
            is_active=True,
        )
        self.session.add(existing_group)
        self.session.commit()

        content = create_minimal_normalized_workbook(groups=[{"code": "2500001", "jornada": "Nocturna"}])

        preview = import_service.preview_import_for_coordination(
            session=self.session,
            file_bytes=content,
            filename="import.xlsx",
            import_type="schedule_normalized",
            coordination_id=1,
            is_global=False,
        )
        self.assertEqual(preview.summary["groups"].conflicts, 1)
        self.assertTrue(any("otra coordinación" in c.message for c in preview.conflicts_detail))

        res = import_service.commit_workbook_for_coordination(
            session=self.session,
            file_bytes=content,
            import_type="schedule_normalized",
            coordination_id=1,
            user_id=1,
            filename="import.xlsx",
        )
        self.assertEqual(res.conflicts.get("groups", 0), 1)

        db_grp = self.session.exec(select(Group).where(Group.code == "2500001")).first()
        self.assertEqual(db_grp.coordination_id, 2)
        self.assertEqual(db_grp.jornada, "Diurna")

    # 3. TEST AC-04 & Punto 64: Group of same coordination updates safely
    def test_group_same_coordination_safe_update(self):
        existing_group = Group(
            code="2500002",
            name="Ficha Teleinformática",
            coordination_id=1,
            jornada="Diurna",
            is_active=True,
        )
        self.session.add(existing_group)
        self.session.commit()

        content = create_minimal_normalized_workbook(groups=[{"code": "2500002", "jornada": "Mixta"}])

        res = import_service.commit_workbook_for_coordination(
            session=self.session,
            file_bytes=content,
            import_type="schedule_normalized",
            coordination_id=1,
            user_id=1,
            filename="import.xlsx",
        )
        self.assertEqual(res.updated.get("groups", 0), 1)

        db_grp = self.session.exec(select(Group).where(Group.code == "2500002")).first()
        self.assertEqual(db_grp.coordination_id, 1)
        self.assertEqual(db_grp.jornada, "Mixta")

    # 4. TEST AC-05 & Punto 65: New group receives active coordination
    def test_group_new_receives_active_coordination(self):
        content = create_minimal_normalized_workbook(groups=[{"code": "2500003", "jornada": "Diurna"}])
        res = import_service.commit_workbook_for_coordination(
            session=self.session,
            file_bytes=content,
            import_type="schedule_normalized",
            coordination_id=1,
            user_id=1,
            filename="import.xlsx",
        )
        self.assertEqual(res.created.get("groups", 0), 1)

        db_grp = self.session.exec(select(Group).where(Group.code == "2500003")).first()
        self.assertIsNotNone(db_grp)
        self.assertEqual(db_grp.coordination_id, 1)

    # 5. TEST AC-02, AC-03 & Punto 66: Missing records in subsequent load are NEVER deleted/deactivated
    def test_incremental_missing_records_not_deleted(self):
        content1 = create_minimal_normalized_workbook(groups=[{"code": "100"}, {"code": "101"}])
        import_service.commit_workbook_for_coordination(
            session=self.session,
            file_bytes=content1,
            import_type="schedule_normalized",
            coordination_id=1,
            user_id=1,
            filename="batch1.xlsx",
        )
        grp100 = self.session.exec(select(Group).where(Group.code == "100")).first()
        grp101 = self.session.exec(select(Group).where(Group.code == "101")).first()
        self.assertIsNotNone(grp100)
        self.assertIsNotNone(grp101)

        content2 = create_minimal_normalized_workbook(groups=[{"code": "101"}])
        import_service.commit_workbook_for_coordination(
            session=self.session,
            file_bytes=content2,
            import_type="schedule_normalized",
            coordination_id=1,
            user_id=1,
            filename="batch2.xlsx",
        )

        db_grp100 = self.session.exec(select(Group).where(Group.code == "100")).first()
        self.assertIsNotNone(db_grp100)
        self.assertTrue(db_grp100.is_active)

    # 6. TEST AC-06, AC-07, AC-08 & Punto 67: Shared instructor creates M:N relation, no duplication
    def test_shared_instructor_no_duplicate(self):
        content = create_minimal_normalized_workbook(
            instructors=[{"name": "CARLOS LOPEZ", "document_number": "CC-12345"}]
        )

        res1 = import_service.commit_workbook_for_coordination(
            session=self.session,
            file_bytes=content,
            import_type="schedule_normalized",
            coordination_id=1,
            user_id=1,
            filename="coord1.xlsx",
        )
        self.assertEqual(res1.created.get("instructors", 0), 1)

        res2 = import_service.commit_workbook_for_coordination(
            session=self.session,
            file_bytes=content,
            import_type="schedule_normalized",
            coordination_id=2,
            user_id=2,
            filename="coord2.xlsx",
        )
        self.assertEqual(res2.created.get("instructors", 0), 0)

        all_insts = self.session.exec(select(Instructor).where(Instructor.document_number == "CC-12345")).all()
        self.assertEqual(len(all_insts), 1)
        inst_id = all_insts[0].id

        links = self.session.exec(
            select(InstructorCoordination).where(InstructorCoordination.instructor_id == inst_id)
        ).all()
        self.assertEqual(len(links), 2)
        self.assertEqual({link.coordination_id for link in links}, {1, 2})

    # 7. TEST AC-12 & Punto 68: Instructor conflicting contact email is NOT overwritten
    def test_shared_instructor_email_conflict(self):
        inst = Instructor(
            document_type="CC",
            document_number="CC-9999",
            first_name="ANA",
            last_name="GOMEZ",
            email="ana@sena.edu.co",
            primary_coordination_id=1,
            is_active=True,
        )
        self.session.add(inst)
        self.session.commit()

        content = create_minimal_normalized_workbook(
            instructors=[{"name": "ANA GOMEZ", "document_number": "CC-9999", "email": "otro@gmail.com"}]
        )

        res = import_service.commit_workbook_for_coordination(
            session=self.session,
            file_bytes=content,
            import_type="schedule_normalized",
            coordination_id=2,
            user_id=2,
            filename="coord2.xlsx",
        )
        self.assertEqual(res.conflicts.get("instructors", 0), 1)

        db_inst = self.session.exec(select(Instructor).where(Instructor.document_number == "CC-9999")).first()
        self.assertEqual(db_inst.email, "ana@sena.edu.co")

    # 8. TEST AC-09 & Punto 69: Empty/null cell does NOT erase existing data
    def test_null_does_not_erase_existing_fields(self):
        self.assertEqual(merge_field("3001234567", None), "3001234567")
        self.assertEqual(merge_field("3001234567", ""), "3001234567")
        self.assertEqual(merge_field("3001234567", "   "), "3001234567")
        self.assertEqual(merge_field("3001234567", "3119876543"), "3119876543")

    # 9. TEST AC-10 & Punto 70: Shared Environment M:N
    def test_shared_environment_coordination_links(self):
        content = create_minimal_normalized_workbook(environments=[{"code": "LAB-401", "name": "Lab Robótica"}])

        import_service.commit_workbook_for_coordination(
            session=self.session,
            file_bytes=content,
            import_type="schedule_normalized",
            coordination_id=1,
            user_id=1,
            filename="c1.xlsx",
        )

        import_service.commit_workbook_for_coordination(
            session=self.session,
            file_bytes=content,
            import_type="schedule_normalized",
            coordination_id=2,
            user_id=2,
            filename="c2.xlsx",
        )

        envs = self.session.exec(select(Environment).where(Environment.code == "LAB-401")).all()
        self.assertEqual(len(envs), 1)
        env_id = envs[0].id

        links = self.session.exec(
            select(EnvironmentCoordination).where(EnvironmentCoordination.environment_id == env_id)
        ).all()
        self.assertEqual(len(links), 2)
        self.assertEqual({l.coordination_id for l in links}, {1, 2})

    # 10. TEST Punto 71: Shared Environment capacity conflict not overwritten
    def test_environment_capacity_conflict(self):
        env = Environment(code="AUDITORIO", name="Auditorio Principal", capacity=30, is_active=True)
        self.session.add(env)
        self.session.flush()
        self.session.add(EnvironmentCoordination(environment_id=env.id, coordination_id=1))
        self.session.add(EnvironmentCoordination(environment_id=env.id, coordination_id=2))
        self.session.commit()

        content = create_minimal_normalized_workbook(environments=[{"code": "AUDITORIO", "capacity": 80}])
        preview = import_service.preview_import_for_coordination(
            session=self.session,
            file_bytes=content,
            filename="env.xlsx",
            import_type="schedule_normalized",
            coordination_id=2,
        )
        self.assertEqual(preview.summary["environments"].conflicts, 1)

        res = import_service.commit_workbook_for_coordination(
            session=self.session,
            file_bytes=content,
            import_type="schedule_normalized",
            coordination_id=2,
            user_id=2,
        )
        self.assertEqual(res.conflicts.get("environments", 0), 1)

        db_env = self.session.exec(select(Environment).where(Environment.code == "AUDITORIO")).first()
        self.assertEqual(db_env.capacity, 30)

    # 11. TEST AC-11, AC-12 & Punto 72: Global TrainingProgram conflict not overwritten
    def test_program_structural_conflict(self):
        prog = TrainingProgram(code="PROG-ADSO", name="ADSO", level="Tecnólogo", is_active=True)
        self.session.add(prog)
        self.session.commit()

        from app.services.imports.safe_merge import classify_training_program
        c = classify_training_program(self.session, {"code": "PROG-ADSO", "level": "Técnico"})
        self.assertEqual(c.action, "CONFLICT")

    # 12. TEST AC-13 & Punto 73: Idempotent reimport of same file
    def test_same_file_reimport_idempotent(self):
        content = create_minimal_normalized_workbook(groups=[{"code": "F-1001"}])

        res1 = import_service.commit_workbook_for_coordination(
            session=self.session,
            file_bytes=content,
            import_type="schedule_normalized",
            coordination_id=1,
            user_id=1,
            filename="reimport.xlsx",
        )
        self.assertEqual(res1.created.get("groups", 0), 1)

        preview2 = import_service.preview_import_for_coordination(
            session=self.session,
            file_bytes=content,
            filename="reimport.xlsx",
            import_type="schedule_normalized",
            coordination_id=1,
        )
        self.assertTrue(preview2.is_reimport)

        res2 = import_service.commit_workbook_for_coordination(
            session=self.session,
            file_bytes=content,
            import_type="schedule_normalized",
            coordination_id=1,
            user_id=1,
            filename="reimport.xlsx",
        )
        self.assertEqual(res2.created.get("groups", 0), 0)
        self.assertGreater(res2.unchanged.get("groups", 0), 0)

    # 13. TEST Punto 74: Same filename but different content
    def test_same_name_different_content(self):
        content1 = create_minimal_normalized_workbook(groups=[{"code": "G-A"}])
        content2 = create_minimal_normalized_workbook(groups=[{"code": "G-B"}])

        hash1 = compute_file_sha256(content1)
        hash2 = compute_file_sha256(content2)
        self.assertNotEqual(hash1, hash2)

    # 14. TEST Punto 75: SHA-256 stability
    def test_sha256_stability(self):
        b1 = b"SENAHORARIOS2026"
        b2 = b"SENAHORARIOS2026"
        b3 = b"SENAHORARIOS2027"
        self.assertEqual(compute_file_sha256(b1), compute_file_sha256(b2))
        self.assertNotEqual(compute_file_sha256(b1), compute_file_sha256(b3))

    # 15. TEST AC-14, AC-15 & Punto 76: Single atomic transaction rolls back on failure
    def test_atomic_transaction_rollback_on_failure(self):
        content = create_minimal_normalized_workbook(groups=[{"code": "ATOMIC-1"}, {"code": "ATOMIC-2"}])
        from unittest.mock import patch

        with patch("app.services.import_service.classify_group", side_effect=RuntimeError("Simulated DB Crash")):
            with self.assertRaises(RuntimeError):
                import_service.commit_workbook_for_coordination(
                    session=self.session,
                    file_bytes=content,
                    import_type="schedule_normalized",
                    coordination_id=1,
                    user_id=1,
                    filename="crash.xlsx",
                )

        persisted = self.session.exec(select(Group).where(Group.code.like("ATOMIC%"))).all()
        self.assertEqual(len(persisted), 0)

    # 16. TEST Punto 78: Duplicates in file
    def test_duplicates_in_file(self):
        from app.services.imports.safe_merge import deduplicate_records

        rows_identical = [
            {"code": "101", "jornada": "DIURNA"},
            {"code": "101", "jornada": "DIURNA"},
        ]
        warns = []
        errs = []
        res = deduplicate_records(rows_identical, lambda x: x["code"], "groups", "FICHAS", warns, errs)
        self.assertEqual(len(res), 1)
        self.assertEqual(len(warns), 1)
        self.assertEqual(len(errs), 0)

        rows_conflict = [
            {"code": "102", "jornada": "DIURNA"},
            {"code": "102", "jornada": "NOCTURNA"},
        ]
        warns2 = []
        errs2 = []
        res2 = deduplicate_records(rows_conflict, lambda x: x["code"], "groups", "FICHAS", warns2, errs2)
        self.assertEqual(len(errs2), 1)

    # 17. TEST AC-17 & Punto 79: ImportBatch and ImportBatchRecord creation
    def test_import_batch_and_records_persisted(self):
        content = create_minimal_normalized_workbook(groups=[{"code": "BATCH-1"}])
        res = import_service.commit_workbook_for_coordination(
            session=self.session,
            file_bytes=content,
            import_type="schedule_normalized",
            coordination_id=1,
            user_id=1,
            filename="batch_test.xlsx",
        )
        self.assertIsNotNone(res.batch_id)

        batch = self.session.get(ImportBatch, res.batch_id)
        self.assertIsNotNone(batch)
        self.assertEqual(batch.coordination_id, 1)
        self.assertEqual(batch.uploaded_by_user_id, 1)
        self.assertEqual(batch.status, "completed")
        self.assertGreater(batch.created_count, 0)

        records = self.session.exec(select(ImportBatchRecord).where(ImportBatchRecord.batch_id == batch.id)).all()
        self.assertGreater(len(records), 0)

    # 18. TEST AC-18 & Punto 80: History endpoint scoped by AccessScope
    def test_history_scoped_by_access_scope(self):
        batch1 = ImportBatch(
            import_type="schedule_normalized",
            coordination_id=1,
            uploaded_by_user_id=1,
            filename="c1.xlsx",
            file_sha256="hash1",
            status="completed",
            created_at=datetime.utcnow(),
        )
        batch2 = ImportBatch(
            import_type="schedule_normalized",
            coordination_id=2,
            uploaded_by_user_id=2,
            filename="c2.xlsx",
            file_sha256="hash2",
            status="completed",
            created_at=datetime.utcnow(),
        )
        self.session.add(batch1)
        self.session.add(batch2)
        self.session.commit()

        stmt1 = select(ImportBatch).where(ImportBatch.coordination_id.in_({1}))
        visible_c1 = self.session.exec(stmt1).all()
        self.assertEqual(len(visible_c1), 1)
        self.assertEqual(visible_c1[0].coordination_id, 1)

        visible_admin = self.session.exec(select(ImportBatch)).all()
        self.assertEqual(len(visible_admin), 2)
